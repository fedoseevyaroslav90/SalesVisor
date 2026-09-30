import os
import time
from datetime import date

import pandas as pd
from fastapi.testclient import TestClient

from salesvisor import parsing as p
from salesvisor.config import Settings
from salesvisor.db import make_engine
from salesvisor.ingest import load_segments, load_svetofor, parse_segments
from salesvisor.queries import list_orders, order_card
from salesvisor.web import create_app


def test_parsing():
    assert p.parse_date("22 июля, 2026") == date(2026, 7, 22)
    assert p.parse_date("8 июня, 2026, 16:44") == date(2026, 6, 8)
    assert p.parse_date("31.12.2025") == date(2025, 12, 31)
    assert p.parse_date(None) is None
    assert p.parse_number("1 814,22") == 1814.22
    assert p.parse_flag("@08@") and not p.parse_flag("@EB@")
    assert p.parse_decade("2Д07", near=date(2026, 7, 22)) == ("2Д07", date(2026, 7, 20))
    assert p.parse_decade("3Д02", near=date(2026, 7, 1)) == ("3Д02", date(2026, 2, 28))
    # Декада декабря рядом с январской датой относится к прошлому году
    assert p.parse_decade("2Д12", near=date(2026, 1, 15)) == ("2Д12", date(2025, 12, 20))
    assert p.parse_bitrix_task("Задача № 313360 Губанов") == "313360"
    assert p.parse_bitrix_task("Громова") is None


def seg_row(order, pos, seg, req, produced="@EB@", shipped="@EB@", invoiced="@EB@", task=None, manager="IVANOVA",
            stock="@EB@", ready="@EB@", transit="@EB@"):
    return {
        "Заказ клиента": order, "Позиция заказа клиента": pos, "Номер отрезка по порядку в позиции": seg,
        "Треб. дата поставки": req, "Создал": manager, "Имя заказчика": "ООО Альфа", "Описание отдела сбыта": "Отдел РФ",
        "Статус Произведен": produced, "Статус На складе": stock, "Статус Готов к отгрузке": ready,
        "Статус Отгружен": shipped, "Статус Отфактурирован": invoiced, "Статус В пути": transit,
        "Длина отдельного отрезка": "2", "ЕИ": "КМ", "Номер задачи в Битрикс": task,
    }


def svet_row(order, pos, first, cur, shift, color):
    return {"Заказ": order, "Позиция": pos, "ПланДатаОтгрузки": "20 октября, 2026", "Дата перевода Z4": "1 сентября, 2026",
            "Первая декада": first, "Текущая декада": cur, "Смещено дней": shift, "Цвет": color,
            "Заказчик": "ООО Альфа", "Отдел продаж": "Отдел РФ", "ГП": "Кабель ОК"}


def test_position_keys_with_thousand_separator():
    segs, pos = parse_segments(pd.DataFrame([seg_row("1200000001", "1 010", "1", "10 октября, 2026")]))
    assert pos[0]["pos"] == "1010"


def test_load_and_track_changes():
    engine = make_engine("sqlite:///:memory:")
    load_segments(engine, pd.DataFrame([
        seg_row("1200000001", "10", "1", "30 сентября, 2026", task="Задача № 313360"),
        seg_row("1200000001", "10", "2", "30 сентября, 2026"),
        seg_row("1200000002", "10", "1", "10 марта, 2026", produced="@08@", shipped="@08@", invoiced="@08@"),
    ]), "day1")
    load_svetofor(engine, pd.DataFrame([svet_row("1200000001", "10", "3Д09", "3Д09", "0", "green")]), "day1")

    # Через день: срок сдвинули, один отрезок произведён, светофор пожелтел
    r = load_segments(engine, pd.DataFrame([
        seg_row("1200000001", "10", "1", "10 октября, 2026", produced="@08@", task="Задача № 313360"),
        seg_row("1200000001", "10", "2", "10 октября, 2026"),
        seg_row("1200000002", "10", "1", "10 марта, 2026", produced="@08@", shipped="@08@", invoiced="@08@"),
    ]), "day2")
    assert r["changes"] == 2  # треб. дата и этап
    load_svetofor(engine, pd.DataFrame([svet_row("1200000001", "10", "3Д09", "2Д10", "20", "yellow")]), "day2")

    card = order_card(engine, "1200000001", today=date(2026, 10, 1))
    p0 = card["positions"][0]
    assert p0["first_decade"] == "3Д09" and p0["current_decade"] == "2Д10" and p0["color"] == "yellow"
    assert p0["stage"] == "В производстве 1/2" and p0["bitrix_task"] == "313360" and p0["manager"] == "IVANOVA"
    assert {c["field"] for c in card["changes"]} == {"required_date", "stage", "current_decade", "color"}

    orders = list_orders(engine, today=date(2026, 10, 1))
    assert [o["order_no"] for o in orders] == ["1200000001"]  # отгруженный заказ закрыт
    assert orders[0]["color"] == "yellow" and not orders[0]["overdue"]
    assert list_orders(engine, today=date(2026, 10, 25))[0]["overdue"] == 1


def test_api_comments_and_upload():
    engine = make_engine("sqlite:///:memory:")
    client = TestClient(create_app(engine, Settings(database_url="sqlite:///:memory:")))
    csv = pd.DataFrame([svet_row("1200000009", "10", "1Д10", "1Д10", "0", "green")]).to_csv(index=False).encode()
    r = client.post("/api/upload", files={"file": ("svetofor.csv", csv, "text/csv")})
    assert r.status_code == 200 and r.json()["source"] == "svetofor"
    assert client.post("/api/orders/1200000009/comments", json={"text": "Звонил клиенту", "author": "Я"}).status_code == 200
    card = client.get("/api/orders/1200000009").json()
    assert card["comments"][0]["text"] == "Звонил клиенту"
    assert client.post("/api/upload", files={"file": ("x.csv", b"a,b\n1,2\n", "text/csv")}).status_code == 400
    assert client.post("/api/sync").status_code == 400  # Metabase не настроен


def test_bitrix_manual_link_and_live():
    import httpx
    from salesvisor.bitrix import Bitrix, post_decade_changes

    engine = make_engine("sqlite:///:memory:")
    load_segments(engine, pd.DataFrame([seg_row("1200000005", "10", "1", "30 сентября, 2026", task="Громова")]), "d1")
    load_svetofor(engine, pd.DataFrame([svet_row("1200000005", "10", "3Д09", "3Д09", "0", "green")]), "d1")
    load_svetofor(engine, pd.DataFrame([svet_row("1200000005", "10", "3Д09", "1Д10", "10", "yellow")]), "d2")

    calls = []

    def handler(request: httpx.Request):
        calls.append((request.url.path, request.content.decode()))
        if request.url.path.endswith("tasks.task.get.json"):
            return httpx.Response(200, json={"result": {"task": {"title": "Заказ 1200000005", "status": "3",
                                                                  "deadline": "2026-10-10T18:00:00+05:00",
                                                                  "responsible": {"name": "Громова Анна"}}}})
        return httpx.Response(200, json={"result": True})

    settings = Settings(database_url="sqlite:///:memory:", bitrix_webhook_url="https://b24.local/rest/1/key")
    app = create_app(engine, settings)
    client = TestClient(app)
    assert client.get("/api/orders/1200000005").json()["bitrix"]["task_id"] is None
    assert client.put("/api/orders/1200000005/bitrix", json={"task_id": "abc"}).status_code == 400
    assert client.put("/api/orders/1200000005/bitrix", json={"task_id": "345579", "author": "Я"}).status_code == 200
    card = client.get("/api/orders/1200000005").json()
    assert card["bitrix"]["task_id"] == "345579" and card["bitrix"]["manual"]["set_by"] == "Я"
    assert client.get("/api/orders").json()[0]["bitrix_task"] == "345579"

    bx = Bitrix(settings, client=httpx.Client(base_url=settings.bitrix_webhook_url + "/", transport=httpx.MockTransport(handler)))
    assert post_decade_changes(engine, bx) == 1
    assert "345579" in calls[-1][1] and "3Д09 → 1Д10" in calls[-1][1]
    assert post_decade_changes(engine, bx) == 0  # второй раз не отправляет

    from salesvisor.bitrix import live_info
    info = live_info(bx, "345579", None)
    assert info["task"]["status"] == "Выполняется" and info["task"]["responsible"] == "Громова Анна"


def test_import_folder_and_portal_guard(tmp_path):
    from salesvisor.sync import run_sync

    engine = make_engine("sqlite:///:memory:")
    inbox = tmp_path / "import"
    inbox.mkdir()
    pd.DataFrame([svet_row("1200000007", "10", "3Д09", "1Д10", "10", "yellow")]).to_csv(inbox / "svetofor.csv", index=False)
    pd.DataFrame([seg_row("1200000007", "10", "1", "30 сентября, 2026")]).to_csv(inbox / "otrezki.csv", index=False)
    (inbox / "readme.csv").write_text("a,b\n1,2\n")
    old = time.time() - 600
    for f in inbox.iterdir():
        os.utime(f, (old, old))
    # файл, который ещё докачивается по SFTP (свежий), в этот проход не берётся
    pd.DataFrame([seg_row("1200000008", "10", "1", "30 сентября, 2026")]).to_csv(inbox / "idet.csv", index=False)
    settings = Settings(database_url="sqlite:///:memory:", import_dir=str(inbox), portal_token="s3cret")

    res = run_sync(engine, settings)
    assert [r["source"] for r in res] == ["segments", "svetofor"]  # отрезки раньше светофора
    assert [f.name for f in inbox.iterdir() if f.is_file()] == ["idet.csv"]
    assert len(list((inbox / "done").iterdir())) == 2 and len(list((inbox / "failed").iterdir())) == 1
    # разобранные файлы старше 30 дней удаляются при следующем проходе
    stale = inbox / "done" / "20250101-000000_old.csv"
    stale.write_text("x")
    ancient = time.time() - 40 * 86400
    os.utime(stale, (ancient, ancient))
    run_sync(engine, settings)
    assert not stale.exists() and len(list((inbox / "done").iterdir())) == 2

    client = TestClient(create_app(engine, settings))
    assert client.get("/api/health").status_code == 200
    assert client.get("/api/orders").status_code == 401
    assert client.get("/", headers={"X-SalesVisor-Token": "wrong"}).status_code == 401
    h = {"X-SalesVisor-Token": "s3cret", "X-SalesVisor-User": "user-ivanova",
         "X-SalesVisor-Person": "%D0%98%D0%B2%D0%B0%D0%BD%D0%BE%D0%B2%D0%B0%20%D0%90."}
    assert client.get("/api/meta", headers=h).json()["portal_user"] == "Иванова А."
    assert client.post("/api/orders/1200000007/comments", json={"text": "Проверка", "author": "кто-то"}, headers=h).status_code == 200
    assert client.get("/api/orders/1200000007", headers=h).json()["comments"][0]["author"] == "Иванова А."
    # автор за порталом — только из заголовков портала: ни тело, ни X-Remote-User его не подменят
    h2 = {"X-SalesVisor-Token": "s3cret", "X-Remote-User": "chuzhoy"}
    client.post("/api/orders/1200000007/comments", json={"text": "Второй", "author": "кто-то"}, headers=h2)
    assert {c["author"] for c in client.get("/api/orders/1200000007", headers=h).json()["comments"]} == {"Иванова А.", "без имени"}

    # роли: смотрящему загрузка и «Обновить сейчас» закрыты, загружающему — открыты
    viewer, editor = {**h, "X-SalesVisor-Role": "viewer"}, {**h, "X-SalesVisor-Role": "editor"}
    csv = pd.DataFrame([svet_row("1200000007", "10", "3Д09", "2Д10", "20", "red")]).to_csv(index=False).encode()
    assert client.get("/api/meta", headers=viewer).json()["can_upload"] is False
    assert client.get("/api/meta", headers=editor).json()["can_upload"] is True
    assert client.post("/api/upload", files={"file": ("s.csv", csv)}, headers=viewer).status_code == 403
    assert client.post("/api/sync", headers=viewer).status_code == 403
    r = client.post("/api/upload", files={"file": ("s.csv", csv)}, headers=editor)
    assert r.status_code == 200 and r.json()["source"] == "svetofor"
    assert client.post("/api/upload", files={"file": ("x.csv", b"a;b\n1;2\n")}, headers=editor).status_code == 400
    assert client.post("/api/upload", files={"file": ("x.xlsx", b"not a zip")}, headers=editor).status_code == 400
    # документация API с внешним CDN за порталом не отдаётся
    assert client.get("/api/docs", headers=h).status_code == 404 and client.get("/openapi.json", headers=h).status_code == 404


def test_dirty_rows_do_not_break_load():
    """Повтор позиции, цвет вне списка, отрезок без номера и слишком длинные строки — загрузка проходит."""
    from salesvisor.db import positions
    from salesvisor.ingest import _fit

    engine = make_engine("sqlite:///:memory:")
    segs = pd.DataFrame([seg_row("1200000021", "10", None, "30 сентября, 2026"), seg_row("1200000021", "10", None, "30 сентября, 2026")])
    r = load_segments(engine, segs, "d1")
    assert r["positions"] == 1 and order_card(engine, "1200000021")["positions"][0]["segments_total"] == 2
    svet = pd.DataFrame([svet_row("1200000021", "10", "3Д09", "3Д09", "0", "green"),
                         svet_row("1200000021", "10", "3Д09", "1Д10", "10", '"><img src=x onerror=alert(1)>')])
    r = load_svetofor(engine, svet, "d1")
    assert r["positions"] == 1
    p0 = order_card(engine, "1200000021")["positions"][0]
    assert p0["current_decade"] == "1Д10" and p0["color"] is None
    row = _fit(positions, [{"first_decade": "3Д12 (2025) перенос", "customer": "x" * 500}])[0]
    assert row["first_decade"] == "3Д12 (2025" and len(row["customer"]) == 300


def test_plan_dispatcher_and_day():
    from salesvisor.ingest import detect_source, load_dispatcher, load_plan
    from salesvisor.queries import day_plan, dispatcher_summary

    engine = make_engine("sqlite:///:memory:")
    load_segments(engine, pd.DataFrame([
        seg_row("1200000011", "10", "1", "10 октября, 2026"),
        seg_row("1200000011", "10", "2", "10 октября, 2026", produced="@08@"),
        seg_row("1200000011", "20", "1", "10 октября, 2026", produced="@08@"),
    ]), "d1")
    plan = pd.DataFrame([
        {"Заказ клиента": "1200000011", "Позиция заказа": "10", "Рабочее место": "SZ-2", "Дата конца": "01.10.2026", "Время конца": "08:10", "Номер ДСЕ": "A"},
        {"Заказ клиента": "1200000011", "Позиция заказа": "10", "Рабочее место": "OEL60-5", "Дата конца": "02.10.2026", "Время конца": "15:23", "Номер ДСЕ": "A"},
        {"Заказ клиента": "1200000011", "Позиция заказа": "20", "Рабочее место": "OEL60-5", "Дата конца": "02.10.2026", "Время конца": "09:00", "Номер ДСЕ": "B",
         "Сообщение": "300018741", "Описание сообщения по качеству": "Слипание оболочки"},
        {"Заказ клиента": "1200009999", "Позиция заказа": "10", "Рабочее место": "OEL60-5", "Дата конца": "02.10.2026", "Время конца": "09:00", "Номер ДСЕ": "C"},
    ])
    assert detect_source(list(plan.columns)) == "plan"
    r = load_plan(engine, plan, "plan.xlsx")
    assert r["positions"] == 2 and r["not_in_orders"] == 1 and r["quality_new"] == 1
    assert load_plan(engine, plan, "plan2.xlsx")["quality_new"] == 0  # повторный прогон не дублирует
    card = order_card(engine, "1200000011")
    p10 = card["positions"][0]
    assert card["quality"][0]["text"] == "Слипание оболочки" and card["quality"][0]["line"] == "OEL60-5"
    assert list_orders(engine)[0]["quality"] == 1
    assert p10["line"] == "OEL60-5" and p10["plan_end_time"] == "15:23" and p10["plan_lines"] == "SZ-2, OEL60-5"

    day = day_plan(engine, date(2026, 10, 2), today=date(2026, 10, 2))
    assert day["summary"]["make"] == 2 and day["summary"]["make_done"] == 1
    assert day["lines"][0]["line"] == "OEL60-5" and day["lines"][0]["make"] == 2
    late = day_plan(engine, date(2026, 10, 3), with_backlog=True, today=date(2026, 10, 3))
    assert [x["pos"] for x in late["positions"] if x["task_make"]] == ["10"]  # произведённая позиция не висит в хвосте

    disp = pd.DataFrame([
        {"Плановые МЗ (Руб)": 100, "ВП": 10, "Готов? (назначена партия - готов, пусто - не готов)": "готов", "ПО": "считать",
         "Признак декады": "27. ОКТЯБРЬ 1декада", "Заказ клиента": "1200000011", "Позиция заказа клиента": "20",
         "Длина отдельного отрезка": 2, "Базовая ЕИ": "КМ", "Количество км волокна": 32, "Партия": "9100000001"},
        {"Плановые МЗ (Руб)": 50, "ВП": 5, "Готов? (назначена партия - готов, пусто - не готов)": "готов", "ПО": "считать",
         "Признак декады": "27. ОКТЯБРЬ 1декада", "Заказ клиента": "1200000011", "Позиция заказа клиента": "10",
         "Длина отдельного отрезка": 1, "Базовая ЕИ": "КМ", "Количество км волокна": 16, "Партия": "9100000002"},
        {"Плановые МЗ (Руб)": 70, "ВП": 7, "Готов? (назначена партия - готов, пусто - не готов)": "не готов", "ПО": "не считать",
         "Признак декады": "27. ОКТЯБРЬ 1декада", "Заказ клиента": "1200000012", "Позиция заказа клиента": "10",
         "Длина отдельного отрезка": 3, "Базовая ЕИ": "КМ", "Количество км волокна": 8, "Партия": None},
    ])
    assert detect_source(list(disp.columns)) == "dispatcher"
    load_dispatcher(engine, disp, "disp.xlsx")
    s = dispatcher_summary(engine, decade="27. ОКТЯБРЬ 1декада")
    dec = s["months"][0]["decades"][0]
    assert dec["km"] == 3 and dec["km_ready"] == 3 and dec["mz"] == 150 and s["not_counted"] == 1
    # Позиция 10: партия назначена, а по отрезкам произведена половина — расхождение
    assert dec["mismatch"] == 1
    assert order_card(engine, "1200000011")["positions"][1]["disp_ready"] == "готов"


def test_load_order_does_not_matter():
    """План и диспетчерский раньше отрезков: строки ждут и применяются, когда позиции появились (30.09.2026)."""
    from salesvisor.db import dispatcher, plan_rows
    from salesvisor.ingest import load_dispatcher, load_plan
    from sqlalchemy import select

    engine = make_engine("sqlite:///:memory:")
    plan = pd.DataFrame([
        {"Заказ клиента": "1200000031", "Позиция заказа": "10", "Рабочее место": "OEL60-5", "Дата конца": "02.10.2026",
         "Время конца": "15:23", "Номер ДСЕ": "A", "Сообщение": "300018800", "Описание сообщения по качеству": "Слипание"},
    ])
    disp = pd.DataFrame([
        {"Плановые МЗ (Руб)": 100, "ВП": 10, "Готов? (назначена партия - готов, пусто - не готов)": "готов", "ПО": "считать",
         "Признак декады": "27. ОКТЯБРЬ 1декада", "Заказ клиента": "1200000031", "Позиция заказа клиента": "10",
         "Длина отдельного отрезка": 2, "Базовая ЕИ": "КМ", "Количество км волокна": 32, "Партия": "9100000031"},
    ])
    r = load_plan(engine, plan, "plan.xlsx")
    assert r["positions"] == 0 and r["not_in_orders"] == 1 and r["quality_new"] == 1
    r = load_dispatcher(engine, disp, "disp.xlsx")
    assert r["positions"] == 0 and r["not_in_orders"] == 1

    r = load_segments(engine, pd.DataFrame([seg_row("1200000031", "10", "1", "10 октября, 2026", produced="@08@")]), "d1")
    assert r["deferred_applied"] == 2
    p = order_card(engine, "1200000031")["positions"][0]
    assert p["line"] == "OEL60-5" and p["plan_end_time"] == "15:23"
    assert p["disp_decade"] == "27. ОКТЯБРЬ 1декада" and p["disp_ready"] == "готов" and p["disp_batch"] == "9100000031"
    # появление строк у новой позиции переносом не считается, повторная загрузка ничего не применяет заново
    assert order_card(engine, "1200000031")["changes"] == []
    assert load_segments(engine, pd.DataFrame([seg_row("1200000031", "10", "1", "10 октября, 2026", produced="@08@")]),
                         "d2")["deferred_applied"] == 0
    with engine.connect() as conn:
        assert all(r.applied for r in conn.execute(select(plan_rows.c.applied)))
        assert all(r.applied for r in conn.execute(select(dispatcher.c.applied)))

    # следующий прогон ППМ без этой позиции не стирает её линию (в ZPP только ближайший горизонт)
    load_plan(engine, plan.assign(**{"Заказ клиента": "1200000099"}), "plan2.xlsx")
    assert order_card(engine, "1200000031")["positions"][0]["line"] == "OEL60-5"


def test_filters_sort_and_export():
    """Отбор на уровне позиции, сортировка и страницы ленты позиций, выгрузка в Excel (30.09.2026)."""
    import io
    from openpyxl import load_workbook
    from salesvisor.queries import Filters, list_orders, list_positions

    engine = make_engine("sqlite:///:memory:")
    load_segments(engine, pd.DataFrame([
        seg_row("1200000041", "10", "1", "10 октября, 2026"),
        seg_row("1200000041", "20", "1", "10 октября, 2026", produced="@08@"),
        seg_row("1200000042", "10", "1", "25 октября, 2026", manager="PETROV"),
    ]), "d1")
    load_svetofor(engine, pd.DataFrame([
        svet_row("1200000041", "10", "1Д10", "3Д10", "20", "red"),
        svet_row("1200000041", "20", "1Д10", "1Д10", "0", "green"),
        svet_row("1200000042", "10", "3Д10", "3Д10", "0", "green"),
    ]).assign(**{"Заказчик": ["=HYPERLINK(\"x\")", "=HYPERLINK(\"x\")", "ООО Бета"]}), "d1")
    today = date(2026, 9, 30)

    # заказ попадает в список, если подошла хоть одна позиция, и считается по подошедшим
    o = list_orders(engine, filters=Filters.from_query(stage="not_made"), today=today)
    assert {x["order_no"]: x["positions"] for x in o} == {"1200000041": 1, "1200000042": 1}
    assert [x["order_no"] for x in list_orders(engine, filters=Filters.from_query(shift_min="10"), today=today)] == ["1200000041"]
    # срок сейчас — текущая декада светофора: 3Д10 = 21–31.10
    f = Filters.from_query(due_from="2026-10-21", due_to="2026-10-31")
    assert {(p["order_no"], p["pos"]) for p in list_positions(engine, filters=f, today=today)["rows"]} == \
        {("1200000041", "10"), ("1200000042", "10")}
    assert list_positions(engine, filters=Filters.from_query(customer="бета"), today=today)["total"] == 1
    assert list_positions(engine, filters=Filters.from_query(color="red"), today=today)["total"] == 1
    # сортировка и страницы
    r = list_positions(engine, sort="-shift_days", limit=2, today=today)
    assert r["total"] == 3 and len(r["rows"]) == 2 and r["rows"][0]["shift_days"] == 20
    r = list_positions(engine, sort="due_date", offset=2, limit=2, today=today)
    assert len(r["rows"]) == 1 and r["rows"][0]["due_date"] == "2026-10-31"

    client = TestClient(create_app(engine, Settings(database_url="sqlite:///:memory:")))
    assert client.get("/api/positions?stage=ready,none&limit=5").json()["total"] == 0
    x = client.get("/api/export.xlsx?view=positions&sort=-shift_days&note=Проба")
    assert x.status_code == 200 and "attachment" in x.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(x.content)).active
    assert ws["A1"].value.startswith("Проба") and ws["A2"].value == "Заказ" and ws.max_row == 5
    assert ws["C3"].value == '=HYPERLINK("x")' and ws["C3"].data_type == "s"  # похожее на формулу — текстом
    assert ws["K3"].number_format == "DD.MM.YYYY"
    x = client.get("/api/export.xlsx?view=orders&color=red")
    assert load_workbook(io.BytesIO(x.content)).active.max_row == 3


def test_reject_goz_inkab_and_soon_tile():
    """Причина отклонения, ГОЗ по префиксу изделия, внутренние заказы ООО «Инкаб», плитка «Срок ≤ 7 дней»."""
    from salesvisor.queries import Filters, list_orders, list_positions

    engine = make_engine("sqlite:///:memory:")
    rows = [seg_row("1200000051", "10", "1", "3 октября, 2026"), seg_row("1200000052", "10", "1", "25 октября, 2026"),
            seg_row("1200000053", "10", "1", "25 октября, 2026")]
    rows[0].update({"Причина отклонения": "Z6", "Описание Причины отклонения": "Прогноз", "Материал": "ГОЗ12 ТОС2-П"})
    rows[1].update({"Имя заказчика": 'ООО "Инкаб"', "Материал": "ОК-БС-01"})
    rows[2].update({"Имя заказчика": 'ООО "ИНКАБ ДАЛЬНИЙ ВОСТОК"', "Причина отклонения": "Z4", "Материал": "ВПК ОКЛ"})
    load_segments(engine, pd.DataFrame(rows), "d1")
    today = date(2026, 9, 30)
    nums = lambda f: sorted(p["order_no"] for p in list_positions(engine, filters=Filters.from_query(**f), today=today)["rows"])
    assert nums({"reject": "Z6"}) == ["1200000051"]
    assert nums({"reject": "-"}) == ["1200000052"]
    assert nums({"reject": "!Z6,Z7"}) == ["1200000052", "1200000053"]
    assert nums({"flags": "goz"}) == ["1200000051", "1200000053"]
    assert nums({"flags": "noinkab"}) == ["1200000051", "1200000053"]  # Дальний Восток не внутренний
    assert [o["order_no"] for o in list_orders(engine, color="soon", today=today)] == ["1200000051"]


def test_multi_value_filters():
    """Выбор нескольких значений: менеджеры, этапы, причины — через «|» (этап и причина — и через запятую)."""
    from salesvisor.queries import Filters, day_plan, list_positions

    engine = make_engine("sqlite:///:memory:")
    load_segments(engine, pd.DataFrame([
        seg_row("1200000061", "10", "1", "3 октября, 2026", manager="IVANOVA"),
        seg_row("1200000062", "10", "1", "3 октября, 2026", manager="PETROV", produced="@08@"),
        seg_row("1200000063", "10", "1", "3 октября, 2026", manager="SIDOROVA"),
    ]), "d1")
    today = date(2026, 9, 30)
    nums = lambda **kw: sorted(p["order_no"] for p in list_positions(
        engine, manager=kw.pop("manager", ""), filters=Filters.from_query(**kw), today=today)["rows"])
    assert nums(manager="IVANOVA|PETROV") == ["1200000061", "1200000062"]
    assert nums(manager="PETROV") == ["1200000062"]
    assert nums(stage="made|not_made") == ["1200000061", "1200000062", "1200000063"]
    assert nums(stage="made,in_prod") == ["1200000062"]
    assert nums(line="-|OEL60-5") == ["1200000061", "1200000062", "1200000063"]
    assert day_plan(engine, date(2026, 10, 3), manager="IVANOVA|SIDOROVA", today=today)["summary"]["positions"] == 0


def test_user_prefs_per_portal_user():
    """Преднастройки хранятся за сотрудником портала: у каждого свои, чужие не видны."""
    engine = make_engine("sqlite:///:memory:")
    client = TestClient(create_app(engine, Settings(database_url="sqlite:///:memory:", portal_token="s3cret")))
    ivan = {"X-SalesVisor-Token": "s3cret", "X-SalesVisor-User": "user-ivanova"}
    petr = {"X-SalesVisor-Token": "s3cret", "X-SalesVisor-User": "user-petrov"}
    assert client.get("/api/prefs", headers=ivan).json() == {"alias": "user-ivanova", "views": [], "last": "", "sap_login": "",
                                                               "hidden_cols": {}}
    assert client.put("/api/prefs", headers=ivan, json={"hidden_cols": {"orders": ["bitrix_task"]}}).status_code == 200
    assert client.get("/api/prefs", headers=ivan).json()["hidden_cols"] == {"orders": ["bitrix_task"]}
    assert client.put("/api/prefs", headers=ivan, json={"views": [{"name": "Мои просрочки", "query": "flags=overdue"}],
                                                         "sap_login": "IVANOVA"}).status_code == 200
    assert client.put("/api/prefs", headers=ivan, json={"last": "line=OEL60-5"}).status_code == 200  # остальное не трогает
    p = client.get("/api/prefs", headers=ivan).json()
    assert p["views"][0]["name"] == "Мои просрочки" and p["sap_login"] == "IVANOVA" and p["last"] == "line=OEL60-5"
    assert client.get("/api/prefs", headers=petr).json()["views"] == []
    assert client.put("/api/prefs", headers=petr, json={"views": [{"name": "x" * 61, "query": ""}]}).status_code == 422
    # без сотрудника от портала преднастроек нет
    assert client.get("/api/prefs", headers={"X-SalesVisor-Token": "s3cret"}).status_code == 400


def test_mp_mz_priority_and_order_expand():
    """МП и МЗ — сумма по отрезкам позиции; приоритет = МП × дни опоздания к первой дате; позиции одного заказа."""
    from salesvisor.queries import Filters, list_orders, list_positions

    engine = make_engine("sqlite:///:memory:")
    rows = [seg_row("1200000071", "10", "1", "10 октября, 2026"), seg_row("1200000071", "10", "2", "10 октября, 2026"),
            seg_row("1200000072", "10", "1", "10 октября, 2026")]
    for r, mp, mz in zip(rows, ("1 000,50", "2 000", "500"), ("300", "700", "100")):
        r.update({"МП (рассчетно) ВАЛ Х КУРС": mp, "МЗ (рассчетно) ВАЛ Х КУРС": mz, "МП (расчетно) по тек. курсу": "999999"})
    load_segments(engine, pd.DataFrame(rows), "d1")
    load_svetofor(engine, pd.DataFrame([svet_row("1200000071", "10", "1Д09", "1Д10", "30", "red"),
                                        svet_row("1200000072", "10", "1Д10", "1Д10", "0", "green")]), "d1")
    today = date(2026, 9, 30)
    p = {x["order_no"]: x for x in list_positions(engine, today=today)["rows"]}
    assert p["1200000071"]["mp_rub"] == 3000.5 and p["1200000071"]["mz_rub"] == 1000  # «ВАЛ Х КУРС», а не «по тек. курсу»
    # первая дата 10.09, срок сейчас 10.10 → опоздание 30 дней; у второй позиции срок = первой дате, но сегодня раньше
    assert p["1200000071"]["days_late"] == 30 and p["1200000071"]["priority"] == round(3000.5 * 30)
    assert p["1200000072"]["days_late"] == 0 and p["1200000072"]["priority"] == 0
    o = {x["order_no"]: x for x in list_orders(engine, today=today)}
    assert o["1200000071"]["mp_rub"] == 3000.5 and o["1200000071"]["priority"] == round(3000.5 * 30)
    assert [x["pos"] for x in list_positions(engine, filters=Filters.from_query(order="1200000072"), today=today)["rows"]] == ["10"]


def test_stats_otd_by_first_date():
    """Отчёт по срокам: OTD к первой дате клиента, распределение по декадам опоздания, глубина просрочки, топ заказчиков."""
    from salesvisor.queries import stats

    engine = make_engine("sqlite:///:memory:")
    shipped = dict(produced="@08@", shipped="@08@", invoiced="@08@")
    rows = [seg_row("1200000081", "10", "1", "10 сентября, 2026", **shipped),   # отгружена 08.09 — в срок
            seg_row("1200000082", "10", "1", "10 сентября, 2026", **shipped),   # отгружена 25.09 — +2 декады
            seg_row("1200000083", "10", "1", "10 сентября, 2026"),              # не отгружена, дата прошла
            seg_row("1200000084", "10", "1", "10 ноября, 2026")]                # в работе
    rows[0]["Факт. дата поставки"] = "8 сентября, 2026"
    rows[1]["Факт. дата поставки"] = "25 сентября, 2026"
    rows[2]["Имя заказчика"] = "ООО Бета"
    load_segments(engine, pd.DataFrame(rows), "d1")
    load_svetofor(engine, pd.DataFrame([svet_row(o, "10", f, f, "0", "green") for o, f in
                                        (("1200000081", "1Д09"), ("1200000082", "1Д09"), ("1200000083", "1Д09"), ("1200000084", "1Д11"))]), "d1")
    s = stats(engine, scope="all", today=date(2026, 9, 30))
    assert s["summary"]["otd_positions"] == {"ok": 1, "of": 3, "pct": 33.3}
    assert s["summary"]["otd_orders"]["ok"] == 1 and s["summary"]["otd_orders"]["of"] == 3
    sep = next(m for m in s["by_month"] if m["month"] == "2026-09")
    assert (sep["в срок"], sep["+2 декады"], sep["не выполнено"], sep["positions"]) == (1, 1, 1, 3)
    assert next(m for m in s["by_month"] if m["month"] == "2026-11")["в работе"] == 1
    assert {d["bucket"]: d["positions"] for d in s["depth"]}["11–30 дн."] == 1  # срок 10.09, сегодня 30.09 — 20 дней
    assert s["top_customers"][0]["customer"] in ("ООО Альфа", "ООО Бета") and s["top_customers"][0]["bad"] >= 1


def test_production_day_starts_at_8():
    """Сутки: календарные по умолчанию; производственные 08:00→08:00 — окончание 03.10 в 05:40 относится к 02.10."""
    from salesvisor.ingest import load_plan
    from salesvisor.queries import day_plan, prod_day

    assert prod_day("2026-10-03", "05:40") == "2026-10-02" and prod_day("2026-10-03", "08:00") == "2026-10-03"
    engine = make_engine("sqlite:///:memory:")
    load_segments(engine, pd.DataFrame([seg_row("1200000091", "10", "1", "10 октября, 2026"),
                                        seg_row("1200000091", "20", "1", "10 октября, 2026")]), "d1")
    load_plan(engine, pd.DataFrame([
        {"Заказ клиента": "1200000091", "Позиция заказа": "10", "Рабочее место": "SZ-2", "Дата конца": "03.10.2026", "Время конца": "05:40", "Номер ДСЕ": "A"},
        {"Заказ клиента": "1200000091", "Позиция заказа": "20", "Рабочее место": "SZ-2", "Дата конца": "03.10.2026", "Время конца": "09:15", "Номер ДСЕ": "B"},
    ]), "plan.xlsx")
    today = date(2026, 10, 1)
    make = lambda d, **kw: [r["pos"] for r in day_plan(engine, d, today=today, **kw)["positions"] if r["task_make"]]
    # по умолчанию — календарные сутки, как в версии передачи и её контрольных цифрах
    assert make(date(2026, 10, 3)) == ["10", "20"] and make(date(2026, 10, 2)) == []
    # производственные сутки с 08:00 — по выбору
    assert make(date(2026, 10, 2), shift_day=True) == ["10"]
    assert make(date(2026, 10, 3), shift_day=True) == ["20"]


def test_feed_events_z6_control_and_late30():
    """Лента: новая позиция и смена причины (Z6 → пусто = принят в производство); отметки Z6 и опоздание > 30 дней."""
    from salesvisor.queries import Filters, list_positions, recent_changes

    engine = make_engine("sqlite:///:memory:")
    r1 = seg_row("1200000101", "10", "1", "10 сентября, 2026")
    r1.update({"Причина отклонения": "Z6", "Описание Причины отклонения": "Прогноз"})
    load_segments(engine, pd.DataFrame([r1]), "d1")
    load_svetofor(engine, pd.DataFrame([svet_row("1200000101", "10", "1Д08", "1Д09", "30", "red")]), "d1")
    today = date(2026, 9, 30)
    flags = lambda f: [p["order_no"] for p in list_positions(engine, filters=Filters.from_query(flags=f), today=today)["rows"]]
    assert flags("z6late") == ["1200000101"]       # декада срока (1Д09) уже началась, а позиция всё ещё Z6
    assert flags("late30") == ["1200000101"]       # первая дата 10.08, сегодня 30.09 — 51 день
    r1 = dict(r1, **{"Причина отклонения": None, "Описание Причины отклонения": None})
    load_segments(engine, pd.DataFrame([r1, seg_row("1200000102", "10", "1", "10 октября, 2026")]), "d2")
    ev = {(c["order_no"], c["field"]): c for c in recent_changes(engine, days=1)}
    assert ev[("1200000101", "reject_code")]["old"] == "Z6" and ev[("1200000101", "reject_code")]["new"] == "(пусто)"
    assert ("1200000102", "new_position") in ev


def test_status_flags_are_not_cumulative():
    """Флаги SAP не накопительные: у отрезка в пути нет «Готов к отгрузке», у отгруженного — «На складе».
    Отрезок в пути — ушёл с завода: позиция готова и закрыта, не просрочена; этап — по самому отстающему отрезку."""
    engine = make_engine("sqlite:///:memory:")
    Y, R = "@08@", "@0A@"
    req = "1 сентября, 2026"
    segs = pd.DataFrame([
        seg_row("1200000091", "10", "1", req, produced=Y, transit=Y),                                   # П··В··
        seg_row("1200000091", "20", "1", req, produced=Y, shipped=Y, invoiced=Y),                       # П···ОФ
        seg_row("1200000091", "20", "2", req, produced=Y, stock=Y, ready=Y),                            # ПСГ···
        seg_row("1200000091", "30", "1", req, produced=Y, stock=Y, ready=R),                            # ПСr··· — не готов
        seg_row("1200000091", "40", "1", req, produced=Y, shipped=Y),                                   # П···О·
        seg_row("1200000091", "40", "2", req),                                                          # ······
    ])
    load_segments(engine, segs, "d1")
    ps = {p["pos"]: p for p in order_card(engine, "1200000091", today=date(2026, 9, 30))["positions"]}
    assert (ps["10"]["stage"], ps["10"]["closed"], ps["10"]["overdue"]) == ("В пути", True, False)
    assert (ps["20"]["stage"], ps["20"]["closed"], ps["20"]["overdue"]) == ("Готов к отгрузке", False, True)
    assert (ps["30"]["stage"], ps["30"]["closed"]) == ("На складе", False)
    assert (ps["40"]["stage"], ps["40"]["closed"]) == ("В производстве 1/2", False)
    o = list_orders(engine, scope="all", today=date(2026, 9, 30))[0]
    assert (o["segments_ready"], o["segments_total"], o["ready_positions"]) == (4, 6, 2)
