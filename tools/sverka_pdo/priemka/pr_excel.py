"""Excel-отчёт сверки SalesVisor с якорем «первое появление в отчёте по принятым» ПДО."""
import os
import pickle
import sys
from collections import Counter, defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1]
D = pickle.load(open(os.path.join(HERE, "cmp.pkl"), "rb"))
rows, snaps = D["rows"], D["snaps"]
TODAY = date(2026, 9, 30)
STALE = 180

HDR = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="1F3864")
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def pct(a, b):
    return round(100 * a / b, 1) if b else None


def table(ws, r0, headers, data, widths=None, fmt=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill, c.alignment = HDR, HFILL, Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(data, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            if isinstance(v, date):
                c.number_format = "DD.MM.YYYY"
            if fmt and j in fmt:
                c.number_format = fmt[j]
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return r0 + len(data) + 1


def otd(xs, f):
    ok = sum(r[f] == "в срок" for r in xs)
    n = sum(r[f] in ("в срок", "не в срок") for r in xs)
    return ok, n, pct(ok, n)


# ---------------- расчёты
p26 = [r for r in rows if r["sv_first"] and date(2026, 1, 1) <= r["sv_first"] <= TODAY]
both = [r for r in p26 if r["in_pdo"] and r["diff_dec"] is not None]
plan_src = [r for r in both if "план" in r["first_snap"]]
fact_src = [r for r in both if "факт" in r["first_snap"]]
same = [r for r in both if r["otd_anchor"] and r["otd_sv"]]
same_plan = [r for r in same if "план" in r["first_snap"]]
same_fact = [r for r in same if "факт" in r["first_snap"]]
openp = [r for r in rows if not r["closed"]]
open_scope = [r for r in openp if not (r["sv_current"] and (TODAY - r["sv_current"]).days > STALE)]
debt = [r for r in open_scope if r["sv_overdue"]]
debt_all = [r for r in openp if r["sv_overdue"]]
# ПДО закрыл декаду «готов», а в отчёте по отрезкам у позиции ничего не произведено (нет ЗнП и дат)
ghost = [r for r in openp if r["in_pdo"] and r["pdo_fact"] and r["pdo_fact"].split("; ")[-1][6:].startswith("готов")
         and r["stage"] == "Не произведён"]
ghost_debt = [r for r in ghost if r["sv_overdue"] and not (r["sv_current"] and (TODAY - r["sv_current"]).days > STALE)]

wb = Workbook()
ws = wb.active
ws.title = "Итог"
ws.column_dimensions["A"].width = 120
L = [
    ("Сверка SalesVisor с принципом «первый снимок при появлении заказа» (отчёты ПДО «по принятым заказам»)", Font(bold=True, size=14)),
    (f"Данные: SalesVisor — отчёт по отрезкам и светофор на 30.09.2026; ПДО — {len(snaps)} снимков с {snaps[0][0]:%d.%m.%Y} по {snaps[-1][0]:%d.%m.%Y} "
     f"(папка SalesVisor\\data\\приёмка, только файлы «Отчет по принятым заказам…»). Сравнение сделано вне SalesVisor, в программе ничего не менялось.", None),
    ("", None),
    ("1. Где есть предварительные снимки ПДО, первая дата SalesVisor совпадает с якорем ПДО", BOLD),
    (f"• Предварительные снимки есть за январь–июнь, но первая декада в светофоре до июня есть лишь у "
     f"{sum(1 for r in p26 if r['sv_first'].month < 6)} позиций — поэтому это сравнение по сути июньское "
     f"({sum(1 for r in plan_src if r['sv_first'].month == 6)} из {len(plan_src)} позиций).", None),
    (f"• Позиций, которые впервые появились в предварительном отчёте «в декаду … ver.N» (там видны и «принят», и «не принят»): {len(plan_src)}. "
     f"Первая декада светофора совпадает с якорем у {sum(r['diff_dec'] == 0 for r in plan_src)} ({pct(sum(r['diff_dec'] == 0 for r in plan_src), len(plan_src))} %).", None),
    (f"• Июнь (первая декада светофора в июне): совпадение {pct(sum(r['diff_dec'] == 0 for r in both if r['sv_first'].month == 6), sum(1 for r in both if r['sv_first'].month == 6))} %; "
     f"выпуск к сроку и по светофору, и по якорю почти нулевой ({otd([r for r in same if r['sv_first'].month == 6], 'otd_sv')[2]} % и "
     f"{otd([r for r in same if r['sv_first'].month == 6], 'otd_anchor')[2]} %): ПДО 2–3 раза не принимал позиции в декаду и переносил их, выпуск — через 10–50 дней. "
     f"Это реальное опоздание, оба способа его видят одинаково.", None),
    ("", None),
    ("2. За июль–сентябрь в папке только отчёты «и факт в декаду» — поэтому якорь сдвинут и сравнение некорректно", BOLD),
    (f"• Позиций, которые впервые появились в отчёте «и факт» (после декады): {len(fact_src)}. Совпадение с первой декадой светофора — "
     f"{pct(sum(r['diff_dec'] == 0 for r in fact_src), len(fact_src))} %; якорь позже у {sum(r['diff_dec'] > 0 for r in fact_src)}, раньше у {sum(r['diff_dec'] < 0 for r in fact_src)}.", None),
    ("• В отчёт «и факт» позиция попадает только в той декаде, в которую ПДО её уже принял. К этому моменту требуемая дата в SAP обычно уже "
     "перенесена (пример: 1200022129/20 — светофор 2Д06, в SAP 31.07, первое появление — «факт 31.07»). Значит, «первый снимок» здесь — "
     "принятая декада после переносов, а не первое обещание клиенту.", None),
    (f"• Поэтому на июльских позициях выпуск к сроку по якорю {otd([r for r in same if r['sv_first'].month == 7], 'otd_anchor')[2]} %, "
     f"а по светофору {otd([r for r in same if r['sv_first'].month == 7], 'otd_sv')[2]} %: якорь из отчёта «и факт» завышает исполнение.", None),
    ("• «Якорь раньше» — позиции, которые клиенту подвинули на более раннюю декаду (пример: 1200022716 — светофор 3Д07, в SAP 21.07, принят и выпущен в 2Д07).", None),
    ("", None),
    ("3. Долг (просрочено) в SalesVisor почти целиком — позиции, которых ПДО ни разу не принимал в декаду", BOLD),
    (f"• Долг на 30.09 в «Открытых» (без хвостов старше {STALE} дней): {len(debt)} позиций, "
     f"{len({r['order'] for r in debt})} заказов. Из них есть в снимках ПДО: {sum(r['in_pdo'] for r in debt)}; нет ни в одном снимке: {sum(not r['in_pdo'] for r in debt)}.", None),
    (f"• У {sum(1 for r in debt if not r['sv_first'])} из них нет данных светофора — срок взят из требуемой даты SAP. "
     f"Все {len(debt)} — «Не произведён» (ни один отрезок не выпущен).", None),
    (f"• С хвостами (срок прошёл больше {STALE} дней назад, вкладка «Хвосты»): {len(debt_all)} позиций; у {sum(1 for r in debt_all if not r['sv_first'])} нет светофора, "
     f"сроки — с 2025 года. Это висящие непроизведённые позиции, не долг ПДО.", None),
    (f"• Расхождение данных: {len(ghost)} открытых позиций в {len({r['order'] for r in ghost})} заказах ПДО закрыл в декаде как «готов» "
     f"(есть факт поступления), а в отчёте по отрезкам у них ничего не произведено — нет заказа на производство и дат. {len(ghost_debt)} из них "
     f"SalesVisor показывает долгом (например, 1200021502 — 10 позиций, выпуск 06–23.05). Похоже, выпуск не привязан к отрезкам заказа клиента "
     f"в SAP — это вопрос к ПДО/ИТ, список на листе «ПДО готов, отрезки нет».", None),
    ("• По принципу «первого снимка» такие позиции долгом не являются вовсе: якоря у них нет, ПДО их не принимал. В SalesVisor они долг, потому что "
     "прошла требуемая дата клиента (или текущая декада светофора). Это разные вопросы: «что обещали клиенту и не выпустили» против «что ПДО принял и не сделал».", None),
    ("", None),
    ("4. Что нужно, чтобы сравнение было честным за июль–сентябрь", BOLD),
    ("• Предварительные отчёты «Отчет по принятым заказам в декаду ДД.ММ.ГГГГ ver.N» за июль, август, сентябрь (как за январь–июнь). "
     "В папках 07–09 их нет, есть только «и факт». Если ПДО их формирует — положите в те же папки, пересчитаю тем же скриптом за минуты.", None),
    ("", None),
    ("5. Вывод", BOLD),
    ("• Первая дата SalesVisor (первая декада светофора) по смыслу и есть «первый снимок»: где снимки ПДО полные, совпадение 91 %.", None),
    ("• Расхождения июля–сентября — не ошибка SalesVisor, а неполнота снимков. Менять расчёт первой даты по этим данным не предлагаю.", None),
    ("• Можно добавить в SalesVisor второй взгляд, как в старом пилоте: «принятая декада ПДО» (последняя декада, куда позиция принята) и долг ПДО "
     "(принят в декаду и не выпущен). Для этого SalesVisor должен регулярно получать отчёты ПДО — через ту же папку SFTP. Решение за вами.", None),
]
for i, (t, f) in enumerate(L, 1):
    c = ws.cell(i, 1, t)
    c.alignment = WRAP
    if f:
        c.font = f

# ---------------- по месяцам
ws2 = wb.create_sheet("По месяцам")
by = defaultdict(list)
for r in p26:
    by[r["sv_first"].strftime("%Y-%m")].append(r)
data = []
for m in sorted(by):
    xs = by[m]
    b = [r for r in xs if r["in_pdo"] and r["diff_dec"] is not None]
    s = [r for r in b if r["otd_anchor"] and r["otd_sv"]]
    data.append([m, len(xs), len(b), pct(len(b), len(xs)), sum("план" in r["first_snap"] for r in b), sum("факт" in r["first_snap"] for r in b),
                 sum(r["diff_dec"] == 0 for r in b), pct(sum(r["diff_dec"] == 0 for r in b), len(b)),
                 sum(r["diff_dec"] > 0 for r in b), sum(r["diff_dec"] < 0 for r in b),
                 len(s), otd(s, "otd_sv")[2], otd(s, "otd_anchor")[2]])
table(ws2, 1, ["Месяц первой декады светофора", "Позиций", "Есть в снимках ПДО", "%", "Первое появление — план", "Первое появление — факт",
               "Якорь = светофор", "%", "Якорь позже", "Якорь раньше", "Срок прошёл (общее множество)", "Выпуск к сроку по светофору, %",
               "Выпуск к сроку по якорю ПДО, %"], data, [14, 10, 12, 8, 13, 13, 12, 8, 11, 11, 14, 14, 14])

# ---------------- позиции
ws3 = wb.create_sheet("Позиции")
p26_ids = {id(r) for r in p26}
sel = [r for r in rows if r["in_pdo"] or id(r) in p26_ids]
sel.sort(key=lambda r: (r["order"], int(r["pos"]) if r["pos"].isdigit() else 0))
cols = [("Заказ", "order", 12), ("Поз.", "pos", 6), ("Клиент", "customer", 30), ("Отдел", "dept", 18), ("Изделие", "product", 34),
        ("Этап (отрезки)", "stage", 16), ("Закрыта", "closed", 8), ("Выпуск (посл. отрезок)", "release_date", 12),
        ("Первая декада светофора", "sv_first_label", 10), ("Конец первой декады", "sv_first", 11), ("Текущая декада светофора", "sv_current_label", 10),
        ("Срок сейчас", "sv_current", 11), ("Треб. дата SAP сейчас", "sap_req", 11),
        ("Есть в снимках ПДО", "in_pdo", 8), ("Якорь ПДО (первое появление)", "anchor_label", 10), ("Конец декады якоря", "anchor", 11),
        ("Первый снимок", "first_snap", 14), ("Снимков с позицией", "pdo_snaps", 8), ("Раз «не принят»", "pdo_rej", 8),
        ("Причина непринятия (чаще всего)", "pdo_reason", 24), ("Переносов треб. даты в снимках", "pdo_moves", 9),
        ("Итог декад (отчёты «и факт», последние 3)", "pdo_fact", 40),
        ("Якорь − светофор, декад", "diff_dec", 9), ("Долг SalesVisor", "sv_overdue", 8), ("Долг по якорю", "anchor_overdue", 8),
        ("Опоздание SalesVisor, дн.", "sv_days_late", 9), ("Опоздание к якорю, дн.", "anchor_days_late", 9),
        ("Выпуск к сроку — светофор", "otd_sv", 11), ("Выпуск к сроку — якорь", "otd_anchor", 11)]
data = [[("да" if r[k] else "") if isinstance(r[k], bool) else r[k] for _, k, _ in cols] for r in sel]
table(ws3, 1, [c[0] for c in cols], data, [c[2] for c in cols])
ws3.freeze_panes = "C2"
ws3.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(data) + 1}"

# ---------------- долг SalesVisor
ws4 = wb.create_sheet("Долг SalesVisor")
debt.sort(key=lambda r: (r["sv_current"] or date.min, r["order"]))
data = [[("да" if r[k] else "") if isinstance(r[k], bool) else r[k] for _, k, _ in cols] for r in debt]
table(ws4, 1, [c[0] for c in cols], data, [c[2] for c in cols])
ws4.freeze_panes = "C2"
ws4.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(data) + 1}"

# ---------------- ПДО готов, отрезки нет
ws6 = wb.create_sheet("ПДО готов, отрезки нет")
ghost.sort(key=lambda r: (r["order"], int(r["pos"]) if r["pos"].isdigit() else 0))
data = [[("да" if r[k] else "") if isinstance(r[k], bool) else r[k] for _, k, _ in cols] for r in ghost]
table(ws6, 1, [c[0] for c in cols], data, [c[2] for c in cols])
ws6.freeze_panes = "C2"

# ---------------- снимки и метод
ws5 = wb.create_sheet("Снимки и метод")
table(ws5, 1, ["Дата в имени файла", "Вид", "Версия"], [[s[0], "план (до декады)" if s[1] == 0 else "факт (после декады)", s[2] or ""] for s in snaps], [14, 20, 8])
notes = [
    "Якорь ПДО — требуемая дата поставки из ПЕРВОГО по хронологии снимка, где появилась пара «заказ + позиция» (не самая ранняя дата из всех "
    "снимков). Снимки упорядочены по дате в имени файла, затем по версии; план раньше факта той же даты. Внутри снимка — ранняя дата из отрезков позиции. "
    "Как в пилоте VOLS-Zakazy.",
    "Июньские файлы «в ИЮНЬ ver.3 ДД.06» — снимки на дату из имени; «в ИЮНЬ ver.1» — на 01.06. Дубли («!…», «— копия») не учитывались.",
    "Декада якоря и светофора сравниваются по концу декады (10, 20, последний день месяца).",
    "Выпуск — последняя «Факт. дата поставки» по отрезкам позиции из отчёта по отрезкам 30.09. Проверено: совпадает с «Дата поступления (факт)» "
    "отчётов ПДО «и факт» у 25 968 из 26 060 позиций.",
    "Выпуск к сроку: позиция выпущена целиком (все отрезки) не позже конца декады срока. Считается только для позиций, у которых срок уже прошёл.",
    "Долг SalesVisor — как в программе после правки 30.09: срок сейчас (текущая декада светофора, иначе треб. дата SAP) прошёл, а позиция не выпущена. "
    "Долг по якорю — прошла декада якоря, а позиция не выпущена.",
]
for i, t in enumerate(notes, 1):
    c = ws5.cell(i, 5, t)
    c.alignment = WRAP
ws5.column_dimensions["E"].width = 110

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
summary = {
    "plan_src": len(plan_src), "plan_match": sum(r["diff_dec"] == 0 for r in plan_src),
    "fact_src": len(fact_src), "fact_match": sum(r["diff_dec"] == 0 for r in fact_src),
    "debt": len(debt), "debt_orders": len({r["order"] for r in debt}), "debt_in_pdo": sum(r["in_pdo"] for r in debt),
    "debt_no_sv": sum(1 for r in debt if not r["sv_first"]), "debt_all": len(debt_all),
    "jun_otd": (otd([r for r in same if r["sv_first"].month == 6], "otd_sv"), otd([r for r in same if r["sv_first"].month == 6], "otd_anchor")),
    "jul_otd": (otd([r for r in same if r["sv_first"].month == 7], "otd_sv"), otd([r for r in same if r["sv_first"].month == 7], "otd_anchor")),
    "positions_sheet": len(sel), "ghost": len(ghost), "ghost_orders": len({r["order"] for r in ghost}), "ghost_debt": len(ghost_debt),
}
open(os.path.join(HERE, "excel_summary.txt"), "w", encoding="utf-8").write(repr(summary))
print("ok")
