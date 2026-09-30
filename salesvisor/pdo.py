"""Путь в производстве (этап 1 инструмента «Путь заказа», 01.10.2026).

Источники: снимки ПДО «Отчёт по принятым заказам в декаду … ver.N» (pdo_plan) и «… и факт в декаду» (pdo_fact),
ежедневный «Отчёт по материалам (Z0+Z4)» (materials), «Загрузка РЦ \\ Сводный версия NNN» (load). Приходят
вложениями задачи Битрикса «ВАЖНЫЕ НОВОСТИ У2» (bitrix_pdo.py) или через папку SFTP.

Что даёт: текущее решение ПДО по каждому отрезку в декаде, дефициты с историей ожидаемых дат, загрузку рабочих мест,
события в журнале изменений (время события — момент публикации файла, поэтому загрузка истории не засоряет свежую
ленту), сводку по позиции для флагов и карточки, «Пульс ПДО» — своевременность и полноту решений ПДО.
Методика и цифры — «Внедрение ИИ\\40_Аналитика и отчёты\\42_Записки и разборы\\Продажи и планирование»."""
from __future__ import annotations

import calendar
import io
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta

from openpyxl import load_workbook
from sqlalchemy import and_, bindparam, delete, func, insert, select, update
from sqlalchemy.engine import Engine

from . import parsing as p
from .brands import decode
from .db import change_log, deficits, pdo_decision, pdo_files, pdo_rows, positions, wc_load, wc_map

KINDS = ("pdo_plan", "pdo_fact", "load", "materials")
KIND_NAMES = {"pdo_plan": "план ПДО (принято / не принято)", "pdo_fact": "итог декады ПДО", "load": "загрузка переделов",
              "materials": "отчёт по материалам"}

RE_DEC = re.compile(r"в\s+декаду\s+(\d{1,2})\.(\d{1,2})\.(\d{2,4})", re.I)
RE_VER = re.compile(r"ver\.?\s*(\d+)|(\d)\s*-?\s*[а-я]{0,3}\s*пул|верси[яи]\s*(\d+)", re.I)
RE_MONTH = re.compile(r"в\s+[А-ЯЁа-яё]+\s+ver\.?\s*(\d+)(?:\s+(\d{1,2})\.(\d{1,2}))?", re.I)
RE_MAT = re.compile(r"на\s+(\d{1,2})\.(\d{1,2})\.(\d{2,4})")
RE_LOAD = re.compile(r"Сводный\s+версия\s+(\d{3})\s*(трудоемкость)?\s+от\s+(\d{1,2})\.(\d{1,2})\.(\d{4})", re.I)
RE_SHEET = re.compile(r"(\d\d)\.(\d\d)\s*-\s*(\d\d)\.(\d\d)")
STATUS_RE = re.compile(r"^(принят[оа]?/непринят[оа]?|готов/неготов)$")


# ---------------------------------------------------------------- общее

def dec_end(d: date | None) -> date | None:
    if not d:
        return None
    if d.day <= 10:
        return d.replace(day=10)
    if d.day <= 20:
        return d.replace(day=20)
    return d.replace(day=calendar.monthrange(d.year, d.month)[1])


def dec_start(d: date) -> date:
    return d.replace(day=1 if d.day <= 10 else 11 if d.day <= 20 else 21)


def _year(y: str) -> int:
    y = int(y)
    return 2000 + y if y < 100 else y


def _d(v) -> date | None:
    d = p.parse_date(v)
    return d if d and date(2015, 1, 1) <= d <= date(2035, 1, 1) else None


def _num(v) -> float | None:
    """Числа SAP бывают с минусом в конце: «11 076,18-»."""
    if isinstance(v, str) and v.strip().endswith("-"):
        n = p.parse_number(v.strip()[:-1])
        return -n if n is not None else None
    return p.parse_number(v)


def _key(v) -> str:
    s = str(v if v is not None else "").strip().replace(" ", "").replace(" ", "")
    return s.split(".")[0]


def _txt(v, n: int) -> str | None:
    s = p.text(v)
    if s is None or s in ("0", "-"):
        return None
    return re.sub(r"\s+", " ", s)[:n]


def kind_by_name(name: str) -> str | None:
    low = (name or "").lower().replace("ё", "е")
    if "отчет по принятым" in low or "отчет по допринятым" in low:
        return "pdo_fact" if "и факт" in low else "pdo_plan"
    if "отчет по материалам" in low:
        return "materials"
    if "сводный версия" in low and "трудоемк" not in low:
        return "load"
    return None


def _wb(data: bytes):
    return load_workbook(io.BytesIO(data), read_only=True, data_only=True)


def _norm(h) -> str:
    return re.sub(r"[\s.]+", "", str(h or "").lower())


def sniff(data: bytes, filename: str) -> str | None:
    """Вид файла ПДО: по имени, для книги без понятного имени — по колонке «Принят / не принят» / «готов/не готов»."""
    if data[:2] != b"PK":
        return None
    kind = kind_by_name(filename)
    if kind:
        return kind
    try:
        wb = _wb(data)
    except Exception:
        return None
    try:
        for ws in wb.worksheets:
            if "отрезк" in ws.title.lower():
                continue
            for i, row in enumerate(ws.iter_rows(max_row=12, values_only=True)):
                cells = [_norm(c) for c in row]
                if "заказклиента" in cells:
                    st = next((c for c in cells if STATUS_RE.match(c)), None)
                    if st:
                        return "pdo_fact" if st.startswith("готов") else "pdo_plan"
        return None
    finally:
        wb.close()


def sort_key(kind: str, name: str) -> tuple:
    """Порядок загрузки истории из папки: по декаде / дате снимка и версии."""
    m = RE_DEC.search(name)
    if m:
        v = RE_VER.search(name[m.end():])
        ver = int(next(g for g in v.groups() if g)) if v else 0
        return (date(_year(m.group(3)), int(m.group(2)), int(m.group(1))), 0 if kind == "pdo_plan" else 1, ver)
    m = RE_MAT.search(name) if kind == "materials" else RE_LOAD.search(name) if kind == "load" else None
    if m:
        g = m.groups()
        d = date(_year(g[2]), int(g[1]), int(g[0])) if kind == "materials" else date(int(g[4]), int(g[3]), int(g[2]))
        return (d, 2, 0)
    return (date.min, 9, 0)


def _event(conn, rows: list[dict]) -> None:
    if not rows:
        return
    for r in rows:
        r["old"] = (r.get("old") or None) and str(r["old"])[:100]
        r["new"] = (r.get("new") or None) and str(r["new"])[:100]
    conn.execute(insert(change_log), rows)


def _file(conn, **kw) -> int:
    kw.setdefault("loaded_at", datetime.now())
    kw["origin"] = (kw.get("origin") or "")[:300]
    return conn.execute(insert(pdo_files).values(**kw)).inserted_primary_key[0]


# ---------------------------------------------------------------- снимки ПДО

def _pdo_meta(name: str) -> tuple[date | None, int]:
    m = RE_DEC.search(name)
    if m:
        v = RE_VER.search(name[m.end():])
        ver = int(next(g for g in v.groups() if g)) if v else 0
        return dec_end(date(_year(m.group(3)), int(m.group(2)), int(m.group(1)))), ver
    m = RE_MONTH.search(name)
    if m:
        return None, int(m.group(1))       # месячный файл («в ИЮНЬ ver.3 08.06») — декада строки по требуемой дате
    return None, 0


def parse_pdo(data: bytes, filename: str) -> tuple[str, list[dict]]:
    """Строки отчёта ПДО по отрезкам. В 2026 — первый лист, в 2025 — лист «исх данные» (внутри книги 2025 ещё и
    полный снимок «Отчёт по отрезкам» — он пропускается)."""
    wb = _wb(data)
    try:
        for ws in wb.worksheets:
            if "отрезк" in ws.title.lower():
                continue
            it = ws.iter_rows(values_only=True)
            hdr = kind = None
            for i, row in enumerate(it):
                cells = [_norm(c) for c in row]
                st = next((c for c in cells if STATUS_RE.match(c)), None)
                if "заказклиента" in cells and st:
                    hdr = {}
                    for j, c in enumerate(cells):
                        hdr.setdefault(c, j)
                    kind = "pdo_fact" if st.startswith("готов") else "pdo_plan"
                    break
                if i > 12:
                    break
            if not hdr:
                continue

            def g(r, *names):
                for n in names:
                    j = hdr.get(_norm(n))
                    if j is not None and j < len(r) and r[j] is not None:
                        return r[j]
                return None

            out = []
            for r in it:
                o, ps = _key(g(r, "Заказ клиента")), _key(g(r, "Позиция заказа"))
                if not o.startswith("12") or not ps:
                    continue
                status = str(g(r, "Принят / не принят", "Принято/не принято", "готов/не готов") or "").strip().lower()
                out.append({
                    "order_no": o, "pos": ps, "seg_no": _key(g(r, "Порядковый номер")) or "1",
                    "req_date": _d(g(r, "Требуемая дата поставки заказа клиента")),
                    "status": status[:40] or None, "accepted": None if not status else not status.startswith("не"),
                    "reason": _txt(g(r, "Причина не принятия", "причина"), 200),
                    "bottleneck": _txt(g(r, "Материал / РЦ"), 200),
                    "wc": _txt(g(r, "Рабочее место"), 40),
                    "move_to": _d(g(r, "Дата, на которую перенести")),
                    "end_date": _d(g(r, "Дата конца")),
                    "qty": _num(g(r, "Кол-во по заказу клиента по отрезку", "Кол-во поступления (план)")),
                    "unit": _txt(g(r, "Базовая ЕИ"), 10),
                    "mz": _num(g(r, "МЗ (из калькуляции) Плановые на отрезок")),
                    "vp": _num(g(r, "Валовая прибыль, руб.", "Валовая прибыль, руб")),
                    "fact_date": _d(g(r, "Дата поступления (факт)")),
                    "fact_qty": _num(g(r, "Количество поступления (факт)", "Кол-во поступления (факт) MES")),
                    "classifier": _txt(g(r, "Классификатор"), 200),
                    "responsible": _txt(g(r, "ответственный"), 100),
                    "product": _txt(g(r, "Наименование материала ГП", "Наименование материала поступления"), 300),
                })
            return kind, out
        raise ValueError("в книге нет листа с колонками «Заказ клиента» и «Принят / не принят» (или «готов/не готов»)")
    finally:
        wb.close()


def _pos_summary(rows: list[dict]) -> dict:
    """Решение по позиции в декаде по её отрезкам: принят / не принят / частично."""
    acc = [r["accepted"] for r in rows if r.get("accepted") is not None]
    rej = next((r for r in rows if r.get("accepted") is False), None)
    st = "принят" if acc and all(acc) else "не принят" if acc and not any(acc) else "частично" if acc else "нет решения"
    return {"status": st, "rej": rej, "n": len(rows), "n_rej": sum(1 for a in acc if not a)}


def _plan_text(decade: date, s: dict, fact: bool) -> str:
    r = s["rej"] or {}
    if fact:
        what = "не готов" if s["status"] in ("не принят", "частично") else "готов"
        why = r.get("classifier") or r.get("reason") or ""
        return f"{what} в {decade:%d.%m}" + (f": {why}" if why and what == "не готов" else "")
    parts = [f"{s['status']} в {decade:%d.%m}" + (f" ({s['n_rej']} из {s['n']} отр.)" if s["status"] == "частично" else "")]
    if r.get("reason"):
        parts.append(r["reason"])
    if r.get("bottleneck"):
        parts.append(r["bottleneck"])
    txt = ": ".join(parts[:1] + [" · ".join(parts[1:])]) if len(parts) > 1 else parts[0]
    if r.get("move_to"):
        txt += f" → {r['move_to']:%d.%m}"
    return txt


def load_pdo(engine: Engine, data: bytes, filename: str, published_at: datetime | None = None,
             attachment_id: str | None = None) -> dict:
    kind, rows = parse_pdo(data, filename)
    decade, ver = _pdo_meta(filename)
    published_at = published_at or datetime.now()
    for r in rows:
        r["decade_end"] = decade or dec_end(r["req_date"]) or dec_end(published_at.date())
    rej = [r for r in rows if r["accepted"] is False]
    with engine.begin() as conn:
        fid = _file(conn, kind=kind, origin=filename, attachment_id=attachment_id, published_at=published_at,
                    decade_end=decade, version=ver, rows=len(rows), rejects=len(rej),
                    no_reason=sum(1 for r in rej if not r["reason"]),
                    no_bottleneck=sum(1 for r in rej if not r["bottleneck"]) if kind == "pdo_plan" else None,
                    no_move=sum(1 for r in rej if not r["move_to"]) if kind == "pdo_plan" else None)
        if rows:
            conn.execute(insert(pdo_rows), [{**r, "file_id": fid, "kind": kind, "version": ver} for r in rows])
        # текущее решение: новая версия (или та же версия, выложенная позже) заменяет прежнюю
        keys = {(r["decade_end"], r["order_no"], r["pos"]) for r in rows}
        old = defaultdict(list)
        orders = sorted({k[1] for k in keys})
        for chunk in (orders[i:i + 500] for i in range(0, len(orders), 500)):
            for d in conn.execute(select(pdo_decision).where(pdo_decision.c.kind == kind, pdo_decision.c.order_no.in_(chunk))):
                if (d.decade_end, d.order_no, d.pos) in keys:
                    old[(d.decade_end, d.order_no, d.pos)].append(dict(d._mapping))
        newer = {k for k in keys if not old.get(k) or all((ver, published_at) >= (o["version"] or 0, o["published_at"] or datetime.min)
                                                         for o in old[k])}
        by_pos = defaultdict(list)
        for r in rows:
            k = (r["decade_end"], r["order_no"], r["pos"])
            if k in newer:
                by_pos[k].append(r)
        if by_pos:
            for (dk, o, ps) in by_pos:
                conn.execute(delete(pdo_decision).where(pdo_decision.c.kind == kind, pdo_decision.c.decade_end == dk,
                                                        pdo_decision.c.order_no == o, pdo_decision.c.pos == ps))
            seen, ins = set(), []
            for k, rs in by_pos.items():
                for r in rs:
                    sk = (k, r["seg_no"])
                    if sk in seen:
                        continue
                    seen.add(sk)
                    ins.append({"kind": kind, "decade_end": k[0], "order_no": k[1], "pos": k[2], "seg_no": r["seg_no"],
                                "file_id": fid, "version": ver, "published_at": published_at,
                                **{f: r[f] for f in ("req_date", "status", "accepted", "reason", "bottleneck", "wc",
                                                     "move_to", "vp", "classifier", "responsible")}})
            conn.execute(insert(pdo_decision), ins)
        # события: смена решения по позиции в декаде; первое появление — только отказ
        events = []
        for k, rs in by_pos.items():
            s_new = _pos_summary(rs)
            s_old = _pos_summary(old[k]) if old.get(k) else None
            if kind == "pdo_fact":
                if s_new["status"] in ("не принят", "частично") and (not s_old or s_old["status"] != s_new["status"]):
                    events.append((k, None, _plan_text(k[0], s_new, True)))
                continue
            if s_old is None:
                if s_new["status"] in ("не принят", "частично"):
                    events.append((k, None, _plan_text(k[0], s_new, False)))
            elif s_old["status"] != s_new["status"]:
                events.append((k, f"{s_old['status']} в {k[0]:%d.%m}", _plan_text(k[0], s_new, False)))
        _event(conn, [{"order_no": k[1], "pos": k[2], "field": kind, "old": o, "new": n, "at": published_at}
                      for k, o, n in events])
        refreshed = refresh_positions(conn, {(k[1], k[2]) for k in keys})
    return {"source": kind, "file": filename, "decade": decade.isoformat() if decade else None, "version": ver,
            "rows": len(rows), "rejects": len(rej), "events": len(events), "positions": refreshed}


# ---------------------------------------------------------------- отчёт по материалам

MAT_COLS = {"Номер продукта": "mat", "Название продукта": "mat_name", "Название группы": "mat_group",
            "Кол-во пост./потребность": "need", "Крайняя дата ожидаемого поступления": "eta", "Заказ клиента": "order",
            "Позиция заказа клиента": "pos", "Обеспеченность": "state", "Треб.дата поставки": "req",
            "Коментарий нормальный": "comment"}


def parse_materials(data: bytes) -> list[dict]:
    wb = _wb(data)
    try:
        ws = next((w for w in wb.worksheets if w.title.strip().lower() == "дефицит"), None)
        if ws is None:
            raise ValueError("в книге нет листа «Дефицит»")
        it = ws.iter_rows(values_only=True)
        hdr = [str(c).strip() if c is not None else "" for c in next(it)]
        ix = {}
        for j, h in enumerate(hdr):
            for k, v in MAT_COLS.items():
                if v not in ix and h.lower().startswith(k.lower()[:22]):
                    ix[v] = j
        if "order" not in ix or "mat" not in ix:
            raise ValueError("на листе «Дефицит» нет колонок «Заказ клиента» и «Номер продукта»")
        out = []
        for r in it:
            g = lambda k: r[ix[k]] if k in ix and ix[k] < len(r) else None  # noqa: E731
            o = _key(g("order"))
            if not o.startswith("12"):
                continue
            out.append({"order_no": o, "pos": _key(g("pos")), "material": _key(g("mat"))[:20],
                        "material_name": _txt(g("mat_name"), 100), "material_group": _txt(g("mat_group"), 60),
                        "need": _num(g("need")), "eta": _d(g("eta")), "state": (_txt(g("state"), 30) or ""),
                        "req_date": _d(g("req")), "comment": _txt(g("comment"), 100)})
        return out
    finally:
        wb.close()


def _threat(eta: date | None, req: date | None) -> bool:
    """Поставка ставит срок под угрозу: даты нет или материал придёт позже, чем за неделю до требуемой даты."""
    return eta is None or req is None or eta > req - timedelta(days=7)


def _is_real(comment: str | None) -> bool:
    c = (comment or "").lower()
    return ("ждём" in c or "ждем" in c) and not re.search(r"\bох\b|склад", c)


def load_materials(engine: Engine, data: bytes, filename: str, published_at: datetime | None = None,
                   attachment_id: str | None = None) -> dict:
    m = RE_MAT.search(filename)
    report = date(_year(m.group(3)), int(m.group(2)), int(m.group(1))) if m else (published_at or datetime.now()).date()
    with engine.connect() as conn:
        last = conn.execute(select(func.max(pdo_files.c.report_date)).where(pdo_files.c.kind == "materials")).scalar()
    if last and report <= last:
        return {"source": "materials", "file": filename, "report_date": report.isoformat(), "skipped": "уже есть отчёт не старше"}
    rows = parse_materials(data)
    cur = {}
    for r in rows:
        if r["state"].lower() != "не обеспечено":
            continue
        k = (r["order_no"], r["pos"], r["material"])
        a = cur.get(k)
        if a is None:
            cur[k] = {**r, "real": _is_real(r["comment"])}
        else:
            if r["eta"] and (not a["eta"] or r["eta"] > a["eta"]):
                a["eta"] = r["eta"]
            a["real"] = a["real"] or _is_real(r["comment"])
    first_run = last is None
    at = datetime.combine(report, datetime.min.time()) + timedelta(hours=9)
    with engine.begin() as conn:
        _file(conn, kind="materials", origin=filename, attachment_id=attachment_id, published_at=published_at or at,
              report_date=report, rows=len(rows), rejects=len(cur))
        existing = {(d.order_no, d.pos, d.material): dict(d._mapping) for d in conn.execute(select(deficits))}
        ins, upd, events = [], [], []
        for k, r in cur.items():
            e = existing.get(k)
            if e is None or not e["active"]:
                row = {"order_no": k[0], "pos": k[1], "material": k[2], "material_name": r["material_name"],
                       "material_group": r["material_group"], "need": r["need"], "state": r["state"], "active": True,
                       "real": r["real"], "eta": r["eta"], "eta_first": r["eta"], "eta_changes": 0, "comment": r["comment"],
                       "req_date": r["req_date"], "first_seen": report, "last_seen": report, "closed_at": None}
                (ins if e is None else upd).append(row)
                if not first_run and r["real"] and _threat(r["eta"], r["req_date"]):
                    events.append((k, "deficit", None, f"{r['material_name'] or k[2]}: ждём поставку"
                                   + (f" {r['eta']:%d.%m}" if r["eta"] else "")))
                continue
            changes = e["eta_changes"] or 0
            if r["eta"] != e["eta"] and e["eta"] and r["eta"]:
                changes += 1
                if not first_run and r["real"] and (r["eta"] - e["eta"]).days >= 3 and _threat(r["eta"], r["req_date"]):
                    events.append((k, "deficit_eta", f"{e['eta']:%d.%m}", f"{r['material_name'] or k[2]}: поставка {r['eta']:%d.%m}"))
            upd.append({**e, "need": r["need"], "state": r["state"], "real": r["real"], "eta": r["eta"] or e["eta"],
                        "eta_changes": changes, "comment": r["comment"], "req_date": r["req_date"], "last_seen": report})
        for k, e in existing.items():
            if e["active"] and k not in cur:
                # закрытие видно в карточке; событием не пишем — иначе лента тонет (до 2 тыс. в день)
                upd.append({**e, "active": False, "closed_at": report})
        if ins:
            conn.execute(insert(deficits), ins)
        if upd:
            for chunk in (upd[i:i + 1000] for i in range(0, len(upd), 1000)):
                conn.execute(delete(deficits).where(
                    and_(deficits.c.order_no == bindparam("o"), deficits.c.pos == bindparam("p"), deficits.c.material == bindparam("m"))),
                    [{"o": u["order_no"], "p": u["pos"], "m": u["material"]} for u in chunk])
                conn.execute(insert(deficits), chunk)
        _event(conn, [{"order_no": k[0], "pos": k[1], "field": f, "old": o, "new": n, "at": at} for k, f, o, n in events])
        touched = {(k[0], k[1]) for k in cur} | {(k[0], k[1]) for k, e in existing.items() if e["active"]}
        refreshed = refresh_positions(conn, touched)
    return {"source": "materials", "file": filename, "report_date": report.isoformat(), "rows": len(rows),
            "not_provided": len(cur), "events": len(events), "positions": refreshed}


# ---------------------------------------------------------------- загрузка переделов

def parse_load(data: bytes, filename: str) -> tuple[str, date, list[dict], dict]:
    m = RE_LOAD.search(filename)
    if not m or m.group(2):
        raise ValueError("имя файла не похоже на «Сводный версия NNN от ДД.ММ.ГГГГ»")
    ver, snap = m.group(1), date(int(m.group(5)), int(m.group(4)), int(m.group(3)))
    wb = _wb(data)
    recs, wpmap = [], {}
    try:
        for ws in wb.worksheets:
            ms = RE_SHEET.fullmatch(ws.title.strip())
            if not ms:
                continue
            try:
                dec = date(snap.year, int(ms.group(4)), int(ms.group(3)))
            except ValueError:
                continue
            rows = list(ws.iter_rows(max_row=6, values_only=True))
            if len(rows) < 6:
                continue
            r3, r5, r6 = rows[2], rows[4], rows[5]
            group = None
            for j, h in enumerate(r3):
                h = str(h or "").strip()
                if not h or j + 1 >= len(r6):
                    continue
                is_group = bool(re.match(r"^(\dД\.|ЛИНИ|МОБИЛЬН|ОКОНЦОВКА)", h))
                is_wp = h.startswith("W") and "_" in h
                if not (is_group or is_wp):
                    continue
                if is_group:
                    group = h.rstrip(" %")
                name = h.rstrip(" %") if is_group else re.sub(r"^W|_\d+_\d+$", "", h)
                if is_wp and group:
                    wpmap[name] = group
                cap = p.parse_number(re.sub(r"[^\d.,]", "", str(r5[j] or ""))) if j < len(r5) else None
                recs.append({"ver": ver, "snap_date": snap, "decade_end": dec, "level": "группа" if is_group else "место",
                             "name": name[:60], "grp": (group or "")[:60], "km": _num(r6[j]), "load": _num(r6[j + 1]), "cap": cap})
    finally:
        wb.close()
    return ver, snap, recs, wpmap


def load_load(engine: Engine, data: bytes, filename: str, published_at: datetime | None = None,
              attachment_id: str | None = None) -> dict:
    ver, snap, recs, wpmap = parse_load(data, filename)
    uniq = {}
    for r in recs:
        uniq[(r["decade_end"], r["level"], r["name"])] = r
    with engine.begin() as conn:
        _file(conn, kind="load", origin=filename, attachment_id=attachment_id, published_at=published_at or datetime.now(),
              report_date=snap, version=int(ver), rows=len(uniq))
        conn.execute(delete(wc_load).where(wc_load.c.ver == ver, wc_load.c.snap_date == snap))
        if uniq:
            conn.execute(insert(wc_load), list(uniq.values()))
        for name, grp in wpmap.items():
            conn.execute(delete(wc_map).where(wc_map.c.name == name))
            conn.execute(insert(wc_map).values(name=name, grp=grp, updated_at=datetime.now()))
        refreshed = refresh_positions(conn, None, only_load=True)
    return {"source": "load", "file": filename, "version": ver, "snapshot": snap.isoformat(), "records": len(uniq),
            "work_places": len(wpmap), "positions": refreshed}


LOADERS = {"pdo_plan": load_pdo, "pdo_fact": load_pdo, "materials": load_materials, "load": load_load}


# ---------------------------------------------------------------- сводка по позиции

def _norm_wc(wc: str | None) -> str | None:
    if not wc:
        return None
    return re.sub(r"^W|_\d+_\d+$", "", wc.strip().split(",")[0].strip()) or None


def _load_lookup(conn) -> dict:
    """Загрузка рабочего места в декаде по последнему снимку версии 010 (твёрдый план)."""
    out, snaps = {}, {}
    for r in conn.execute(select(wc_load.c.snap_date, wc_load.c.decade_end, wc_load.c.name, wc_load.c.load)
                          .where(wc_load.c.ver == "010", wc_load.c.level == "место")):
        k = (r.name, r.decade_end)
        if k not in snaps or r.snap_date > snaps[k]:
            snaps[k], out[k] = r.snap_date, r.load
    return out


def refresh_positions(conn, keys: set[tuple[str, str]] | None, only_load: bool = False) -> int:
    """Пересчитать поля позиции из снимков ПДО, дефицитов и загрузки. keys=None — все позиции."""
    q = select(positions.c.order_no, positions.c.pos, positions.c.product, positions.c.brand,
               positions.c.pdo_last_wc, positions.c.pdo_last_decade)
    pos = [r for r in conn.execute(q) if keys is None or (r.order_no, r.pos) in keys]
    if not pos:
        return 0
    loads = _load_lookup(conn)
    groups = {r.name: r.grp for r in conn.execute(select(wc_map))}
    if only_load:
        params = []
        for r in pos:
            wc = _norm_wc(r.pdo_last_wc)
            params.append({"k_o": r.order_no, "k_p": r.pos, "wc_group": groups.get(wc) if wc else None,
                           "wc_load": loads.get((wc, r.pdo_last_decade)) if wc else None})
        _bulk_update(conn, params, ("wc_group", "wc_load"))
        return len(params)
    orders = sorted({r.order_no for r in pos})
    dec = defaultdict(lambda: defaultdict(list))
    first = {}
    defs = defaultdict(list)
    for chunk in (orders[i:i + 500] for i in range(0, len(orders), 500)):
        for d in conn.execute(select(pdo_decision).where(pdo_decision.c.order_no.in_(chunk))):
            dec[(d.order_no, d.pos)][(d.kind, d.decade_end)].append(dict(d._mapping))
        for r in conn.execute(select(pdo_rows.c.order_no, pdo_rows.c.pos, pdo_rows.c.req_date, pdo_rows.c.decade_end,
                                     pdo_files.c.published_at, pdo_files.c.id)
                              .join(pdo_files, pdo_files.c.id == pdo_rows.c.file_id)
                              .where(pdo_rows.c.order_no.in_(chunk), pdo_rows.c.kind == "pdo_plan")
                              .order_by(pdo_files.c.published_at, pdo_files.c.id)):
            first.setdefault((r.order_no, r.pos), dec_end(r.req_date) or r.decade_end)
        for d in conn.execute(select(deficits).where(deficits.c.order_no.in_(chunk), deficits.c.active.is_(True))):
            defs[(d.order_no, d.pos)].append(d)
    params = []
    for r in pos:
        k = (r.order_no, r.pos)
        dd = decode(r.product) if r.product else None
        row = {"k_o": k[0], "k_p": k[1], "brand": dd["brand"] if dd else r.brand,
               "product_group": dd["group"] if dd else None, "pdo_first_date": first.get(k)}
        plans = sorted((kd for kd in dec[k] if kd[0] == "pdo_plan"), key=lambda x: x[1])
        facts = sorted((kd for kd in dec[k] if kd[0] == "pdo_fact"), key=lambda x: x[1])
        rej_decades = [kd for kd in plans if any(x["accepted"] is False for x in dec[k][kd])]
        row["pdo_rejects"] = len(rej_decades) if plans else None
        if plans:
            last = plans[-1]
            s = _pos_summary(dec[k][last])
            src = s["rej"] or dec[k][last][0]
            row.update(pdo_last_decade=last[1], pdo_last_status=s["status"], pdo_last_reason=src.get("reason"),
                       pdo_last_bottleneck=src.get("bottleneck"), pdo_last_wc=src.get("wc"), pdo_last_move=src.get("move_to"))
        else:
            row.update(pdo_last_decade=None, pdo_last_status=None, pdo_last_reason=None, pdo_last_bottleneck=None,
                       pdo_last_wc=None, pdo_last_move=None)
        if facts:
            s = _pos_summary(dec[k][facts[-1]])
            src = s["rej"] or {}
            row.update(pdo_fact_status="готов" if s["status"] == "принят" else "не готов" if s["status"] == "не принят"
                       else s["status"], pdo_fact_reason=src.get("classifier") or src.get("reason"))
        else:
            row.update(pdo_fact_status=None, pdo_fact_reason=None)
        wc = _norm_wc(row["pdo_last_wc"])
        row["wc_group"] = groups.get(wc) if wc else None
        row["wc_load"] = loads.get((wc, row["pdo_last_decade"])) if wc else None
        ds = defs.get(k, [])
        row["deficit_active"] = len(ds) or None
        row["deficit_real"] = any(d.real for d in ds) if ds else None
        row["deficit_eta"] = max((d.eta for d in ds if d.eta), default=None)
        row["deficit_materials"] = ", ".join(sorted({d.material_name or d.material for d in ds}))[:300] or None
        params.append(row)
    _bulk_update(conn, params, ("brand", "product_group", "pdo_first_date", "pdo_rejects", "pdo_last_decade", "pdo_last_status",
                                "pdo_last_reason", "pdo_last_bottleneck", "pdo_last_wc", "pdo_last_move", "pdo_fact_status",
                                "pdo_fact_reason", "wc_group", "wc_load", "deficit_active", "deficit_real", "deficit_eta",
                                "deficit_materials"))
    return len(params)


def _bulk_update(conn, params: list[dict], fields: tuple) -> None:
    if not params:
        return
    lengths = {c.name: c.type.length for c in positions.columns if getattr(c.type, "length", None)}
    for r in params:
        for f in fields:
            if isinstance(r.get(f), str) and f in lengths:
                r[f] = r[f][:lengths[f]]
    stmt = (update(positions).where(positions.c.order_no == bindparam("k_o"), positions.c.pos == bindparam("k_p"))
            .values({f: bindparam(f) for f in fields}))
    conn.execute(stmt, params)


def refresh_brands(conn) -> None:
    """Марка и группа продукции у всех позиций (после загрузки отрезков и светофора)."""
    params = []
    for r in conn.execute(select(positions.c.order_no, positions.c.pos, positions.c.product)):
        if r.product:
            d = decode(r.product)
            params.append({"k_o": r.order_no, "k_p": r.pos, "brand": d["brand"], "product_group": d["group"]})
    _bulk_update(conn, params, ("brand", "product_group"))


# ---------------------------------------------------------------- карточка позиции

def card(conn, order_no: str) -> dict:
    """Путь в производстве по позициям заказа: решения ПДО по декадам, итоги декад, дефициты."""
    loads = _load_lookup(conn)
    groups = {r.name: r.grp for r in conn.execute(select(wc_map))}
    dec = defaultdict(lambda: defaultdict(list))
    for d in conn.execute(select(pdo_decision).where(pdo_decision.c.order_no == order_no)):
        dec[d.pos][(d.kind, d.decade_end)].append(dict(d._mapping))
    out = defaultdict(lambda: {"plan": [], "fact": [], "deficits": []})
    for pos_, kd in dec.items():
        for (kind, decade), rs in sorted(kd.items(), key=lambda x: x[0][1]):
            s = _pos_summary(rs)
            src = s["rej"] or rs[0]
            wc = _norm_wc(src.get("wc"))
            item = {"decade": decade.isoformat(), "status": s["status"], "segments": s["n"], "rejected": s["n_rej"],
                    "reason": src.get("reason"), "bottleneck": src.get("bottleneck"), "wc": src.get("wc"),
                    "wc_group": groups.get(wc) if wc else None, "wc_load": loads.get((wc, decade)) if wc else None,
                    "move_to": src.get("move_to").isoformat() if src.get("move_to") else None,
                    "version": src.get("version"), "published_at": src["published_at"].isoformat() if src.get("published_at") else None,
                    "classifier": src.get("classifier"), "responsible": src.get("responsible")}
            out[pos_]["plan" if kind == "pdo_plan" else "fact"].append(item)
    for d in conn.execute(select(deficits).where(deficits.c.order_no == order_no)
                          .order_by(deficits.c.active.desc(), deficits.c.last_seen.desc())):
        out[d.pos]["deficits"].append({k: (v.isoformat() if isinstance(v, date) else v) for k, v in d._mapping.items()})
    return dict(out)


# ---------------------------------------------------------------- пульс ПДО

SIGNAL_NO_PLAN_DAYS = 3      # предварительного отчёта нет за столько дней до начала декады
SIGNAL_NO_FACT_DAYS = 3      # итога нет через столько дней после конца декады
SIGNAL_NO_REASON = 0.2       # доля отказов без причины


def pulse(engine: Engine, today: date | None = None, decades: int = 12) -> dict:
    today = today or date.today()
    with engine.connect() as conn:
        files = [dict(r._mapping) for r in conn.execute(select(pdo_files).where(pdo_files.c.kind.in_(("pdo_plan", "pdo_fact"))))]
        ld = conn.execute(select(func.max(wc_load.c.snap_date)).where(wc_load.c.ver == "010")).scalar()
        load_rows = [dict(r._mapping) for r in conn.execute(select(wc_load).where(wc_load.c.ver == "010", wc_load.c.snap_date == ld,
                                                                                  wc_load.c.level == "группа"))] if ld else []
        last_mat = conn.execute(select(func.max(pdo_files.c.report_date)).where(pdo_files.c.kind == "materials")).scalar()
        unconfirmed = Counter()
        lk = _load_lookup(conn)
        for r in conn.execute(select(pdo_decision.c.decade_end, pdo_decision.c.reason, pdo_decision.c.wc)
                              .where(pdo_decision.c.kind == "pdo_plan", pdo_decision.c.accepted.is_(False))):
            if "загрузк" in (r.reason or "").lower():
                load_ = lk.get((_norm_wc(r.wc), r.decade_end))
                if load_ is not None and load_ < 90:
                    unconfirmed[r.decade_end] += 1
    by = defaultdict(lambda: {"plan": [], "fact": []})
    for f in files:
        if f["decade_end"]:
            by[f["decade_end"]]["plan" if f["kind"] == "pdo_plan" else "fact"].append(f)
    cur = dec_end(today)
    wanted = sorted({d for d in by if d <= dec_end(today + timedelta(days=12))}, reverse=True)[:decades]
    nxt = dec_end(cur + timedelta(days=1))
    for d in (cur, nxt):
        if d not in wanted:
            wanted.append(d)
    rows, signals = [], []
    for d in sorted(set(wanted), reverse=True):
        pl = sorted(by[d]["plan"], key=lambda f: (f["published_at"] or datetime.min, f["id"]))
        fa = sorted(by[d]["fact"], key=lambda f: (f["published_at"] or datetime.min, f["id"]))
        start = dec_start(d)
        first_pub = pl[0]["published_at"].date() if pl and pl[0]["published_at"] else None
        last_plan = max(pl, key=lambda f: (f["version"] or 0, f["published_at"] or datetime.min)) if pl else None
        fact_pub = fa[0]["published_at"].date() if fa and fa[0]["published_at"] else None
        rej = last_plan["rejects"] if last_plan else None
        share = lambda n: round(100 * n / rej) if rej and n is not None else None  # noqa: E731
        row = {"decade": d.isoformat(), "start": start.isoformat(), "plan_published": first_pub.isoformat() if first_pub else None,
               "lead_days": (start - first_pub).days if first_pub else None, "versions": len(pl),
               "fact_published": fact_pub.isoformat() if fact_pub else None,
               "fact_lag_days": (fact_pub - d).days if fact_pub else None,
               "rows": last_plan["rows"] if last_plan else None, "rejects": rej,
               "no_reason_pct": share(last_plan["no_reason"]) if last_plan else None,
               "no_bottleneck_pct": share(last_plan["no_bottleneck"]) if last_plan else None,
               "no_move_pct": share(last_plan["no_move"]) if last_plan else None,
               "load_unconfirmed": unconfirmed.get(d, 0)}
        rows.append(row)
        if not pl and (start - today).days <= SIGNAL_NO_PLAN_DAYS and d >= cur:
            signals.append({"level": "warn", "decade": d.isoformat(),
                            "text": f"Декада {d:%d.%m}: предварительного отчёта «принято / не принято» нет"
                                    + (" — декада уже идёт" if start <= today else f" за {max(0, (start - today).days)} дн. до начала")})
        elif first_pub and first_pub > start:
            signals.append({"level": "info", "decade": d.isoformat(),
                            "text": f"Декада {d:%d.%m}: первый отчёт вышел через {(first_pub - start).days} дн. после начала декады"})
        if not fa and d < today and (today - d).days >= SIGNAL_NO_FACT_DAYS:
            signals.append({"level": "warn", "decade": d.isoformat(),
                            "text": f"Декада {d:%d.%m}: итога «и факт» нет ({(today - d).days} дн. после конца)"})
        if rej and last_plan["no_reason"] is not None and last_plan["no_reason"] / rej > SIGNAL_NO_REASON:
            signals.append({"level": "warn", "decade": d.isoformat(),
                            "text": f"Декада {d:%d.%m}: {row['no_reason_pct']} % отказов без причины ({last_plan['no_reason']} из {rej})"})
    load = defaultdict(dict)
    decs = sorted({r["decade_end"] for r in load_rows if r["decade_end"] >= (ld or today) - timedelta(days=10)})[:7]
    for r in load_rows:
        if r["decade_end"] in decs and r["load"] is not None:
            load[r["name"]][r["decade_end"].isoformat()] = round(r["load"])
    load_table = [{"group": g, **v} for g, v in sorted(load.items(), key=lambda x: -max(x[1].values() or [0])) if max(v.values() or [0]) >= 20]
    return {"today": today.isoformat(), "decades": rows, "signals": signals,
            "load": {"snapshot": ld.isoformat() if ld else None, "decades": [d.isoformat() for d in decs], "rows": load_table},
            "materials_report": last_mat.isoformat() if last_mat else None}
