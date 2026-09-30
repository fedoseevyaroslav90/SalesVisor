"""Выгрузка текущего отбора в Excel: заказы или позиции — с теми же фильтрами и сортировкой, что на экране."""
from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLOR_RU = {"red": "красный", "yellow": "жёлтый", "green": "зелёный"}

# (заголовок, ширина, как взять значение)
ORDER_COLUMNS = [
    ("Заказ", 13, lambda o: o["order_no"]),
    ("Клиент", 36, lambda o: o.get("customer")),
    ("Отдел продаж", 24, lambda o: o.get("sales_dept")),
    ("Менеджер", 14, lambda o: o.get("manager")),
    ("Светофор", 10, lambda o: COLOR_RU.get(o.get("color"), "")),
    ("Позиций", 9, lambda o: o.get("positions")),
    ("Красных", 9, lambda o: o.get("red")),
    ("Жёлтых", 9, lambda o: o.get("yellow")),
    ("Просрочено позиций", 11, lambda o: o.get("overdue")),
    ("Первая декада", 10, lambda o: o.get("first_decade")),
    ("Первая дата клиента", 12, lambda o: o.get("earliest_first")),
    ("Текущая декада", 10, lambda o: o.get("current_decade")),
    ("Ближайший срок", 12, lambda o: o.get("nearest_due")),
    ("Макс. смещение, дн.", 11, lambda o: o.get("max_shift")),
    ("Отрезков готово", 10, lambda o: o.get("segments_ready")),
    ("Отрезков всего", 10, lambda o: o.get("segments_total")),
    ("Позиций готово", 10, lambda o: o.get("ready_positions")),
    ("Сумма заказа, руб", 14, lambda o: o.get("amount_rub")),
    ("МП, руб", 14, lambda o: o.get("mp_rub")),
    ("МЗ, руб", 14, lambda o: o.get("mz_rub")),
    ("Опоздание к первой дате, дн.", 12, lambda o: o.get("days_late")),
    ("Приоритет (МП × дни)", 16, lambda o: o.get("priority")),
    ("Линии", 18, lambda o: ", ".join(o.get("lines") or [])),
    ("Этапы", 28, lambda o: ", ".join(o.get("stages") or [])),
    ("Несоответствий", 10, lambda o: o.get("quality")),
    ("Комментариев", 10, lambda o: o.get("comments")),
    ("Задача Битрикс24", 12, lambda o: o.get("bitrix_task")),
    ("Сделка Битрикс24", 12, lambda o: o.get("bitrix_deal")),
]

POSITION_COLUMNS = [
    ("Заказ", 13, lambda p: p["order_no"]),
    ("Поз.", 6, lambda p: p["pos"]),
    ("Клиент", 32, lambda p: p.get("customer")),
    ("Отдел продаж", 22, lambda p: p.get("sales_dept")),
    ("Менеджер", 13, lambda p: p.get("manager")),
    ("Изделие", 40, lambda p: p.get("product")),
    ("Светофор", 10, lambda p: COLOR_RU.get(p.get("color"), "")),
    ("Первая декада", 10, lambda p: p.get("first_decade")),
    ("Первая дата клиента", 12, lambda p: p.get("first_decade_end")),
    ("Текущая декада", 10, lambda p: p.get("current_decade")),
    ("Срок сейчас", 12, lambda p: p.get("due_date")),
    ("Смещение, дн.", 10, lambda p: p.get("shift_days")),
    ("Просрочено", 10, lambda p: "да" if p.get("overdue") else ""),
    ("Треб. дата поставки", 12, lambda p: p.get("required_date")),
    ("План отгрузки", 12, lambda p: p.get("plan_ship_date") or p.get("invoice_plan_date")),
    ("Линия", 12, lambda p: p.get("line")),
    ("План. окончание", 12, lambda p: p.get("plan_end_date")),
    ("Время окончания", 9, lambda p: p.get("plan_end_time")),
    ("Этап", 22, lambda p: p.get("stage")),
    ("Причина отклонения", 10, lambda p: p.get("reject_code")),
    ("Отрезков готово", 10, lambda p: p.get("segments_ready")),
    ("Отрезков всего", 10, lambda p: p.get("segments_total")),
    ("Длина", 10, lambda p: p.get("length_plan")),
    ("Сумма, руб", 14, lambda p: p.get("amount_rub")),
    ("МП, руб", 14, lambda p: p.get("mp_rub")),
    ("МЗ, руб", 14, lambda p: p.get("mz_rub")),
    ("Опоздание к первой дате, дн.", 12, lambda p: p.get("days_late")),
    ("Приоритет (МП × дни)", 16, lambda p: p.get("priority")),
    ("ЕИ", 6, lambda p: p.get("unit")),
    ("Декада в диспетчерском", 22, lambda p: p.get("disp_decade")),
    ("Готовность в диспетчерском", 16, lambda p: p.get("disp_ready")),
    ("Несоответствий", 10, lambda p: p.get("quality")),
    ("Комментариев к заказу", 10, lambda p: p.get("comments")),
    ("Задача Битрикс24", 12, lambda p: p.get("bitrix_task")),
    ("Сделка Битрикс24", 12, lambda p: p.get("bitrix_deal")),
]


def _value(v):
    """Даты — датами Excel (ISO-строки из выборок превращаются обратно в date)."""
    if isinstance(v, str) and len(v) == 10 and v[4] == "-" and v[7] == "-":
        try:
            return date.fromisoformat(v)
        except ValueError:
            return v
    return v


def build_xlsx(rows: list[dict], kind: str, note: str) -> bytes:
    """Книга с одним листом: строка-подпись отбора, шапка с автофильтром и закреплением, строки."""
    columns = ORDER_COLUMNS if kind == "orders" else POSITION_COLUMNS
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Заказы" if kind == "orders" else "Позиции")
    for i, (_, width, _) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A3"
    last = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A2:{last}{len(rows) + 2}"

    title = WriteOnlyCell(ws, value=note)
    title.font = Font(italic=True, color="555555")
    ws.append([title])
    head = []
    for name, _, _ in columns:
        c = WriteOnlyCell(ws, value=name)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEFC")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        head.append(c)
    ws.append(head)
    for r in rows:
        line = []
        for _, _, get in columns:
            v = _value(get(r))
            if isinstance(v, date):
                c = WriteOnlyCell(ws, value=v)
                c.number_format = "DD.MM.YYYY"
                line.append(c)
            elif isinstance(v, str) and v.startswith("="):
                # строка из выгрузки SAP, похожая на формулу, остаётся текстом (openpyxl записал бы её формулой)
                c = WriteOnlyCell(ws, value=v)
                c.data_type = "s"
                line.append(c)
            else:
                line.append(v)
        ws.append(line)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def file_name(kind: str) -> str:
    return f"SalesVisor_{'заказы' if kind == 'orders' else 'позиции'}_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
