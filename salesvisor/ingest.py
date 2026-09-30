"""Загрузка выгрузок светофора и отчёта по отрезкам в базу с журналом изменений."""
from __future__ import annotations

import hashlib
import io
import re
import threading
from collections import defaultdict
from contextlib import contextmanager
from datetime import date, datetime

import pandas as pd
from sqlalchemy import Table, bindparam, delete, insert, select, text, update
from sqlalchemy.engine import Engine

from . import parsing as p
from .db import change_log, dispatcher, positions, quality_msgs, segments, snapshots

SVETOFOR_REQUIRED = ["Заказ", "Позиция", "Первая декада", "Текущая декада"]
SEGMENTS_REQUIRED = ["Заказ клиента", "Позиция заказа клиента", "Номер отрезка по порядку в позиции", "Треб. дата поставки"]

# План производства и диспетчерский отчёт SAP: колонки ищем по нормализованному заголовку
# (регистр, точки и пробелы не важны). Наборы взяты из рабочих выгрузок, которые читал пилот
PLAN_SIGNATURE = {"заказ клиента", "позиция заказа", "рабочее место", "дата конца", "номер дсе"}
DISPATCHER_SIGNATURE = {"заказ клиента", "позиция заказа клиента", "признак декады", "плановые мз руб"}

# Какие поля позиции отслеживаем между выгрузками
TRACKED_SVETOFOR = ["current_decade", "color", "plan_ship_date"]
TRACKED_SEGMENTS = ["required_date", "stage"]
TRACKED_PLAN = ["line", "plan_end_date"]
TRACKED_DISPATCHER = ["disp_decade", "disp_ready"]

COLORS = {"red", "yellow", "green"}

STAGES = [
    # (поле-счётчик, название этапа, когда все отрезки дошли до него)
    ("segments_invoiced", "Отфактурирован"),
    ("segments_shipped", "Отгружен"),
    ("segments_in_transit", "В пути"),
    ("segments_ready", "Готов к отгрузке"),
    ("segments_stock", "На складе"),
    ("segments_produced", "Произведён"),
]


def read_table(data: bytes, filename: str) -> pd.DataFrame:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        df = pd.read_excel(io.BytesIO(data), sheet_name=pick_sheet(data), dtype=object)
    else:
        text = _decode(data)
        df = pd.read_csv(io.StringIO(text), dtype=str, sep=None, engine="python")
    df.columns = [str(c).strip() for c in df.columns]
    return df


def _sheet_headers(data: bytes) -> list[tuple[str, list]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [(ws.title, [c for c in next(ws.iter_rows(max_row=1, values_only=True), ()) if c is not None])
                for ws in wb.worksheets]
    finally:
        wb.close()


def pick_sheet(data: bytes, want: str | None = None):
    """Лист с данными: в книге диспетчерского первым идёт сводный лист, а строки — на листе «данные»."""
    for title, headers in _sheet_headers(data):
        src = detect_source([str(h).strip() for h in headers])
        if src and (want is None or src == want):
            return title
    return 0


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1251")


def _key(v) -> str | None:
    """Номер заказа или позиции: «10», 10, 10.0 → «10»."""
    s = p.text(v)
    if s is None:
        return None
    s = s.replace("\u00a0", "").replace(" ", "")  # в CSV из Metabase «1 010»
    try:
        f = float(s.replace(",", "."))
        if f.is_integer():
            return str(int(f))
    except ValueError:
        pass
    return s


def _check_columns(df: pd.DataFrame, required: list[str], what: str) -> None:
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"В выгрузке «{what}» нет колонок: {', '.join(missing)}")


def _fmt(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, date):
        return v.isoformat()
    return str(v)


# ---------------------------------------------------------------- светофор

def parse_svetofor(df: pd.DataFrame) -> list[dict]:
    _check_columns(df, SVETOFOR_REQUIRED, "Светофор")
    rows: dict[tuple[str, str], dict] = {}  # позиция дважды в выгрузке — берём последнюю строку
    for r in df.to_dict("records"):
        order_no, pos = _key(r.get("Заказ")), _key(r.get("Позиция"))
        if not order_no or not pos:
            continue
        plan = p.parse_date(r.get("ПланДатаОтгрузки"))
        z4 = p.parse_date(r.get("Дата перевода Z4"))
        first, first_end = p.parse_decade(r.get("Первая декада"), near=z4 or plan)
        cur, cur_end = p.parse_decade(r.get("Текущая декада"), near=plan)
        color = (p.text(r.get("Цвет")) or "").lower()
        rows[(order_no, pos)] = {
            "order_no": order_no, "pos": pos,
            "customer": p.text(r.get("Заказчик")),
            "sales_dept": p.text(r.get("Отдел продаж")),
            "product": p.text(r.get("ГП")),
            "plan_ship_date": plan, "z4_date": z4,
            "first_decade": first, "first_decade_end": first_end,
            "current_decade": cur, "current_decade_end": cur_end,
            "shift_days": p.parse_int(r.get("Смещено дней")),
            "color": color if color in COLORS else None,  # значение уходит в класс CSS интерфейса
            "in_svetofor": True,
        }
    return list(rows.values())


# ---------------------------------------------------------------- отрезки

def parse_segments(df: pd.DataFrame) -> tuple[list[dict], list[dict]]:
    """Возвращает (отрезки, сводку по позициям)."""
    _check_columns(df, SEGMENTS_REQUIRED, "Отчёт по отрезкам")
    product_col = "Материал.1" if "Материал.1" in df.columns else "Материал"
    amount_col = next((c for c in df.columns if c.startswith("СУММА ЗАКАЗА, РУБ. (СУММА ВАЛ")), None)

    segs: dict[tuple[str, str, str], dict] = {}
    agg: dict[tuple[str, str], dict] = {}
    for r in df.to_dict("records"):
        order_no, pos = _key(r.get("Заказ клиента")), _key(r.get("Позиция заказа клиента"))
        if not order_no or not pos:
            continue
        seg = {
            "order_no": order_no, "pos": pos,
            "seg_no": _key(r.get("Номер отрезка по порядку в позиции"))
            or str(agg[(order_no, pos)]["segments_total"] + 1 if (order_no, pos) in agg else 1),
            "length": p.parse_number(r.get("Длина отдельного отрезка")),
            "unit": p.text(r.get("ЕИ")),
            "fact_length": p.parse_number(r.get("Фактическая длина")),
            "produced": p.parse_flag(r.get("Статус Произведен")),
            "stock": p.parse_flag(r.get("Статус На складе")),
            "ready": p.parse_flag(r.get("Статус Готов к отгрузке")),
            "in_transit": p.parse_flag(r.get("Статус В пути")),
            "shipped": p.parse_flag(r.get("Статус Отгружен")),
            "invoiced": p.parse_flag(r.get("Статус Отфактурирован")),
            "fact_ship_date": p.parse_date(r.get("Факт. дата поставки")),
            "prod_order": _key(r.get("Заказ на производство")),
            "warehouse": p.text(r.get("Наименование склада отгрузки")),
        }
        segs[(order_no, pos, seg["seg_no"])] = seg

        a = agg.get((order_no, pos))
        if a is None:
            raw_task = p.text(r.get("Номер задачи в Битрикс"))
            a = agg[(order_no, pos)] = {
                "order_no": order_no, "pos": pos,
                "customer": p.text(r.get("Имя заказчика")),
                "sales_dept": p.text(r.get("Описание отдела сбыта")),
                "manager": p.text(r.get("Создал")),
                "product": p.text(r.get(product_col)),
                "required_date": p.parse_date(r.get("Треб. дата поставки")),
                "invoice_plan_date": p.parse_date(r.get("План дата отгрузки (фактуры)")),
                "to_production_date": p.parse_date(r.get("Дата передачи в производство")),
                "reject_code": p.text(r.get("Причина отклонения")),
                "reject_text": p.text(r.get("Описание Причины отклонения")),
                "unit": seg["unit"],
                "bitrix_raw": raw_task,
                "bitrix_task": p.parse_bitrix_task(raw_task),
                "segments_total": 0, "segments_produced": 0, "segments_stock": 0, "segments_ready": 0,
                "segments_in_transit": 0, "segments_shipped": 0, "segments_invoiced": 0,
                "length_plan": 0.0, "amount_rub": 0.0, "last_fact_ship_date": None,
            }
        a["segments_total"] += 1
        for flag in ("produced", "stock", "ready", "in_transit", "shipped", "invoiced"):
            a[f"segments_{flag}"] += int(seg[flag])
        a["length_plan"] += seg["length"] or 0
        if amount_col:
            a["amount_rub"] += p.parse_number(r.get(amount_col)) or 0
        if seg["fact_ship_date"] and (not a["last_fact_ship_date"] or seg["fact_ship_date"] > a["last_fact_ship_date"]):
            a["last_fact_ship_date"] = seg["fact_ship_date"]

    for a in agg.values():
        a["stage"] = stage_of(a)
        total = a["segments_total"]
        a["closed"] = total > 0 and (a["segments_shipped"] == total or a["segments_invoiced"] == total)
    return list(segs.values()), list(agg.values())


def stage_of(a: dict) -> str:
    total = a.get("segments_total") or 0
    if not total:
        return "Нет отрезков"
    for field, name in STAGES:
        if (a.get(field) or 0) >= total:
            return name
    produced = a.get("segments_produced") or 0
    if produced:
        return f"В производстве {produced}/{total}"
    return "Не произведён"


# ---------------------------------------------------------------- план производства

def _cols(df: pd.DataFrame) -> dict[str, str]:
    """Нормализованный заголовок → исходное имя колонки (первое вхождение)."""
    out: dict[str, str] = {}
    for c in df.columns:
        out.setdefault(p.norm_header(c), c)
    return out


def _get(r: dict, cols: dict[str, str], *names: str):
    for n in names:
        if n in cols:
            v = r.get(cols[n])
            if not p.is_blank(v):
                return v
    return None


def parse_plan(df: pd.DataFrame) -> tuple[list[dict], list[dict]]:
    """ZPP context: план выпуска готовой продукции по позициям, меняется после каждого прогона ППМ.
    Строк по позиции может быть несколько (разные линии, старые плановые заказы). Линия и окончание — у строки
    с самым поздним плановым концом; план и факт MES складываются. Возвращает (позиции, сообщения о качестве)."""
    cols = _cols(df)
    missing = PLAN_SIGNATURE - set(cols)
    if missing:
        raise ValueError(f"В плане производства нет колонок: {', '.join(sorted(missing))}")
    agg: dict[tuple[str, str], dict] = {}
    msgs: dict[str, dict] = {}
    for r in df.to_dict("records"):
        order_no = _key(_get(r, cols, "заказ клиента"))
        pos = _key(_get(r, cols, "позиция заказа", "позиция заказа клиента"))
        if not order_no or not pos:
            continue
        line = p.text(_get(r, cols, "рабочее место"))
        end_d = p.parse_date(_get(r, cols, "дата конца"))
        end_t = p.parse_time(_get(r, cols, "время конца")) or ""
        a = agg.setdefault((order_no, pos), {"order_no": order_no, "pos": pos, "_lines": [], "_best": None,
                                             "dse": None, "plan_msg": None, "plan_qty": 0.0, "plan_fact_qty": 0.0})
        if line and line not in a["_lines"]:
            a["_lines"].append(line)
        a["dse"] = a["dse"] or _key(_get(r, cols, "номер дсе"))
        a["plan_qty"] += p.parse_number(_get(r, cols, "кол во поступления план")) or 0
        a["plan_fact_qty"] += p.parse_number(_get(r, cols, "кол во поступления факт mes", "кол во поступления факт")) or 0
        msg_no = _key(_get(r, cols, "сообщение"))
        msg_text = p.text(_get(r, cols, "описание сообщения по качеству"))
        if msg_no or msg_text:
            a["plan_msg"] = a["plan_msg"] or (msg_text or f"сообщение {msg_no}")[:200]
            key = msg_no or "h" + hashlib.md5(f"{order_no}|{pos}|{msg_text}".encode()).hexdigest()[:19]
            msgs[key] = {"msg_no": key, "order_no": order_no, "pos": pos, "line": line, "text": (msg_text or "")[:300] or None,
                         "product": p.text(_get(r, cols, "наименование материала гп", "наименование материала поступления")),
                         "plan_end_date": end_d}
        k = (end_d or date.min, end_t)
        if end_d and (a["_best"] is None or k > a["_best"][0]):
            a["_best"] = (k, line)
    rows = []
    for a in agg.values():
        best = a.pop("_best")
        lines = a.pop("_lines")
        a["line"] = best[1] if best and best[1] else (lines[-1] if lines else None)
        a["plan_lines"] = ", ".join(lines)[:200] or None
        a["plan_end_date"] = best[0][0] if best else None
        a["plan_end_time"] = (best[0][1] or None) if best else None
        rows.append(a)
    return rows, list(msgs.values())


# ---------------------------------------------------------------- диспетчерский отчёт

def _decade_no(label: str | None) -> int | None:
    m = re.match(r"\s*(\d+)\.", label or "")
    return int(m.group(1)) if m else None


def _row_product(r: dict, columns) -> str | None:
    """В диспетчерском две колонки «Материал»: наименование и код. Берём ту, где есть буквы."""
    for c in columns:
        if p.norm_header(c).startswith("материал"):
            v = p.text(r.get(c))
            if v and re.search("[A-Za-zА-Яа-я]", v):
                return v
    return None


def parse_dispatcher(df: pd.DataFrame) -> list[dict]:
    """Лист «данные» диспетчерского отчёта: строка на отрезок. Итоги считаются как в листе «отчет (итог)»:
    план — все отрезки декады с ПО «считать», факт — из них те, где назначена партия («готов»).
    ГП км — длина отрезков в КМ, ГП шт — количество в ШТ, ОВ — км волокна."""
    cols = _cols(df)
    missing = DISPATCHER_SIGNATURE - set(cols)
    if missing:
        raise ValueError(f"В диспетчерском отчёте нет колонок: {', '.join(sorted(missing))}")
    ready_col = next((c for n, c in cols.items() if n.startswith("готов")), None)
    agg: dict[tuple[str, str], dict] = {}
    for r in df.to_dict("records"):
        order_no = _key(_get(r, cols, "заказ клиента"))
        pos = _key(_get(r, cols, "позиция заказа клиента"))
        if not order_no or not pos:
            continue
        decade = p.text(_get(r, cols, "признак декады"))
        ready = (p.text(r.get(ready_col)) or "").lower() == "готов" if ready_col else False
        unit = (p.text(_get(r, cols, "базовая еи", "еи")) or "").upper()
        qty = p.parse_number(_get(r, cols, "длина отдельного отрезка")) or 0
        vals = {"km": qty if unit == "КМ" else 0, "pcs": qty if unit == "ШТ" else 0,
                "ov_km": p.parse_number(_get(r, cols, "количество км волокна")) or 0,
                "mz": p.parse_number(_get(r, cols, "плановые мз руб")) or 0,
                "vp": p.parse_number(_get(r, cols, "вп")) or 0}
        a = agg.get((order_no, pos))
        if a is None:
            a = agg[(order_no, pos)] = {
                "order_no": order_no, "pos": pos, "decade": decade, "decade_no": _decade_no(decade),
                "customer": p.text(_get(r, cols, "имя заказчика")),
                "product": p.text(_row_product(r, df.columns)),
                "counted": (p.text(_get(r, cols, "по")) or "").lower() == "считать",
                "batch": None, "segs": 0, "segs_ready": 0,
                **{k: 0.0 for k in vals}, **{f"{k}_ready": 0.0 for k in vals}}
        a["segs"] += 1
        a["segs_ready"] += int(ready)
        a["batch"] = a["batch"] or _key(_get(r, cols, "партия"))
        for k, v in vals.items():
            a[k] += v
            if ready:
                a[f"{k}_ready"] += v
    return list(agg.values())


def parse_dispatcher_end_dates(df: pd.DataFrame) -> list[dict]:
    """Лист 1S0D диспетчерского: плановое окончание производства позиции (дата и время конца)."""
    cols = _cols(df)
    rows: dict[tuple[str, str], dict] = {}
    for r in df.to_dict("records"):
        order_no = _key(_get(r, cols, "заказ клиента"))
        pos = _key(_get(r, cols, "позиция заказа", "позиция заказа клиента"))
        end_d = p.parse_date(_get(r, cols, "дата конца"))
        if not order_no or not pos or not end_d:
            continue
        end_t = p.parse_time(_get(r, cols, "время конца"))
        cur = rows.get((order_no, pos))
        if cur is None or (end_d, end_t or "") > (cur["plan_end_date"], cur["plan_end_time"] or ""):
            rows[(order_no, pos)] = {"order_no": order_no, "pos": pos, "plan_end_date": end_d, "plan_end_time": end_t}
    return list(rows.values())


# ---------------------------------------------------------------- запись в базу

_LENGTHS: dict[str, dict[str, int]] = {}


def _fit(table: Table, rows: list[dict]) -> list[dict]:
    """Строки не длиннее колонки. На PostgreSQL слишком длинное значение роняет всю загрузку
    (StringDataRightTruncation), на SQLite длина не проверяется — поэтому обрезаем до записи и до сравнения."""
    lengths = _LENGTHS.get(table.name)
    if lengths is None:
        lengths = _LENGTHS[table.name] = {c.name: c.type.length for c in table.columns
                                          if getattr(c.type, "length", None)}
    for r in rows:
        for k, n in lengths.items():
            v = r.get(k)
            if isinstance(v, str) and len(v) > n:
                r[k] = v[:n]
    return rows


# Одна загрузка за раз: веб (ручная загрузка, «Обновить сейчас») и контейнер синхронизации пишут в одну базу.
# Две загрузки отрезков подряд давали дубль ключа (DELETE второй ждёт первую и не видит её строк),
# светофор вместе с отрезками — взаимоблокировку. На PostgreSQL — рекомендательная блокировка на сеанс,
# на SQLite (локально, один процесс) — блокировка потока.
_LOCK_KEY = 5_417_320_930
_local_lock = threading.Lock()


class LoadBusy(RuntimeError):
    pass


@contextmanager
def load_lock(engine: Engine, wait: bool = True):
    if engine.dialect.name != "postgresql":
        if not _local_lock.acquire(blocking=wait):
            raise LoadBusy("Идёт другая загрузка — повторите через минуту")
        try:
            yield
        finally:
            _local_lock.release()
        return
    with engine.connect() as conn:
        if wait:
            conn.execute(text("SELECT pg_advisory_lock(:k)"), {"k": _LOCK_KEY})
        elif not conn.execute(text("SELECT pg_try_advisory_lock(:k)"), {"k": _LOCK_KEY}).scalar():
            conn.rollback()
            raise LoadBusy("Идёт другая загрузка — повторите через минуту")
        conn.commit()  # блокировка сеансовая: переживает конец транзакции, соединение не висит «в транзакции»
        try:
            yield
        finally:
            conn.execute(text("SELECT pg_advisory_unlock(:k)"), {"k": _LOCK_KEY})
            conn.commit()


def _load_existing(conn, keys: set[tuple[str, str]], fields: list[str]) -> dict:
    cols = [positions.c.order_no, positions.c.pos] + [positions.c[f] for f in fields]
    out = {}
    for row in conn.execute(select(*cols)):
        k = (row[0], row[1])
        if k in keys:
            out[k] = dict(zip(fields, row[2:]))
    return out


def _upsert(conn, rows: list[dict], tracked: list[str], snapshot_id: int) -> dict:
    rows = _fit(positions, [dict(r) for r in rows])
    keys = {(r["order_no"], r["pos"]) for r in rows}
    existing = _load_existing(conn, keys, tracked)
    now = datetime.now()
    to_insert, to_update, changes = [], [], []
    for r in rows:
        k = (r["order_no"], r["pos"])
        r = {**r, "updated_at": now}
        if k not in existing:
            to_insert.append({**r, "first_seen_at": now})
            continue
        old = existing[k]
        for f in tracked:
            if f in r and r[f] is not None and old.get(f) != r[f] and old.get(f) is not None:
                changes.append({"order_no": k[0], "pos": k[1], "field": f, "old": _fmt(old[f]),
                                "new": _fmt(r[f]), "snapshot_id": snapshot_id, "at": now})
        to_update.append(r)

    if to_insert:
        # В одной пачке у всех строк должен быть одинаковый набор полей
        cols = sorted({c for r in to_insert for c in r})
        conn.execute(insert(positions), [{c: r.get(c) for c in cols} for r in to_insert])
    if to_update:
        fields = sorted({c for r in to_update for c in r} - {"order_no", "pos"})
        stmt = (update(positions)
                .where(positions.c.order_no == bindparam("k_order"), positions.c.pos == bindparam("k_pos"))
                .values({f: bindparam(f) for f in fields}))
        conn.execute(stmt, [{"k_order": r["order_no"], "k_pos": r["pos"], **{f: r.get(f) for f in fields}} for r in to_update])
    if changes:
        conn.execute(insert(change_log), _fit(change_log, changes))
    return {"new": len(to_insert), "updated": len(to_update), "changes": len(changes)}


def _snapshot(conn, source: str, origin: str, rows: int) -> int:
    res = conn.execute(insert(snapshots).values(source=source, origin=(origin or "")[:200], rows=rows, loaded_at=datetime.now()))
    return res.inserted_primary_key[0]


def load_svetofor(engine: Engine, df: pd.DataFrame, origin: str) -> dict:
    rows = parse_svetofor(df)
    with engine.begin() as conn:
        sid = _snapshot(conn, "svetofor", origin, len(df))
        # Позиции, которых нет в свежем светофоре, помечаем, но не удаляем
        conn.execute(update(positions).values(in_svetofor=False))
        stats = _upsert(conn, rows, TRACKED_SVETOFOR, sid)
        _fill_order_managers(conn)
    return {"source": "svetofor", "rows": len(df), "positions": len(rows), **stats}


def load_segments(engine: Engine, df: pd.DataFrame, origin: str) -> dict:
    segs, pos_rows = parse_segments(df)
    segs = _fit(segments, segs)
    with engine.begin() as conn:
        sid = _snapshot(conn, "segments", origin, len(df))
        conn.execute(delete(segments))
        for i in range(0, len(segs), 5000):
            conn.execute(insert(segments), segs[i:i + 5000])
        stats = _upsert(conn, pos_rows, TRACKED_SEGMENTS, sid)
        _fill_order_managers(conn)
    return {"source": "segments", "rows": len(df), "positions": len(pos_rows), **stats}


def load_plan(engine: Engine, df: pd.DataFrame, origin: str) -> dict:
    # Позиции, которых нет в свежем плане, не трогаем: в ZPP context только ближайший горизонт.
    # Сдвиги окончания и смена линии между прогонами ППМ попадают в журнал изменений
    rows, msgs = parse_plan(df)
    with engine.begin() as conn:
        sid = _snapshot(conn, "plan", origin, len(df))
        known, unknown = _split_known(conn, rows)
        stats = _upsert(conn, known, TRACKED_PLAN, sid)
        new_msgs = _save_quality(conn, msgs, sid)
    return {"source": "plan", "rows": len(df), "positions": len(known), "not_in_orders": unknown,
            "quality_msgs": len(msgs), "quality_new": new_msgs, **stats}


def _save_quality(conn, msgs: list[dict], snapshot_id: int) -> int:
    msgs = _fit(quality_msgs, msgs)
    have = {r[0] for r in conn.execute(select(quality_msgs.c.msg_no))}
    new = [m for m in msgs if m["msg_no"] not in have]
    now = datetime.now()
    if new:
        conn.execute(insert(quality_msgs), [{**m, "first_seen_at": now} for m in new])
        # Новое несоответствие — событие в журнале позиции, если позиция нам известна
        known = {(o, ps) for o, ps in conn.execute(select(positions.c.order_no, positions.c.pos))}
        log = [{"order_no": m["order_no"], "pos": m["pos"], "field": "quality", "old": None,
                "new": (f"{m['text'] or 'без описания'} ({m['line'] or 'линия ?'})")[:100], "snapshot_id": snapshot_id, "at": now}
               for m in new if (m["order_no"], m["pos"]) in known]
        if log:
            conn.execute(insert(change_log), _fit(change_log, log))
    old = [{"k_msg": m["msg_no"], "text": m["text"], "line": m["line"], "plan_end_date": m["plan_end_date"]}
           for m in msgs if m["msg_no"] in have]
    if old:
        conn.execute(update(quality_msgs).where(quality_msgs.c.msg_no == bindparam("k_msg"))
                     .values(text=bindparam("text"), line=bindparam("line"), plan_end_date=bindparam("plan_end_date")), old)
    return len(new)


def load_dispatcher(engine: Engine, df: pd.DataFrame, origin: str, end_dates: pd.DataFrame | None = None) -> dict:
    rows = _fit(dispatcher, parse_dispatcher(df))
    with engine.begin() as conn:
        sid = _snapshot(conn, "dispatcher", origin, len(df))
        conn.execute(delete(dispatcher))
        for i in range(0, len(rows), 5000):
            conn.execute(insert(dispatcher), [{k: v for k, v in r.items() if k != "batch"} for r in rows[i:i + 5000]])
        pos_rows = [{"order_no": r["order_no"], "pos": r["pos"], "disp_decade": r["decade"], "disp_counted": r["counted"],
                     "disp_batch": r["batch"],
                     "disp_ready": "готов" if r["segs_ready"] == r["segs"] else "не готов" if not r["segs_ready"]
                     else f"готово {r['segs_ready']} из {r['segs']}"} for r in rows]
        known, unknown = _split_known(conn, pos_rows)
        stats = _upsert(conn, known, TRACKED_DISPATCHER, sid)
        if end_dates is not None:
            ends, _ = _split_known(conn, parse_dispatcher_end_dates(end_dates))
            _upsert(conn, ends, [], sid)
    return {"source": "dispatcher", "rows": len(df), "positions": len(known), "not_in_orders": unknown, **stats}


def _split_known(conn, rows: list[dict]) -> tuple[list[dict], int]:
    """План и диспетчерский дополняют позиции из отчёта по отрезкам и светофора, новых заказов не создают."""
    have = {(o, ps) for o, ps in conn.execute(select(positions.c.order_no, positions.c.pos))}
    known = [r for r in rows if (r["order_no"], r["pos"]) in have]
    return known, len(rows) - len(known)


LOADERS = {"segments": load_segments, "svetofor": load_svetofor, "plan": load_plan, "dispatcher": load_dispatcher}
# Порядок загрузки из папки: сначала отрезки (менеджер, этапы), затем остальное
SOURCE_ORDER = {"segments": 0, "svetofor": 1, "plan": 2, "dispatcher": 3}


def _fill_order_managers(conn) -> None:
    """В светофоре нет менеджера: берём его с других позиций того же заказа."""
    by_order: dict[str, str] = {}
    missing = defaultdict(list)
    for order_no, pos, manager in conn.execute(select(positions.c.order_no, positions.c.pos, positions.c.manager)):
        if manager:
            by_order.setdefault(order_no, manager)
        else:
            missing[order_no].append(pos)
    params = [{"k_order": o, "k_pos": ps, "manager": by_order[o]} for o, lst in missing.items() if o in by_order for ps in lst]
    if params:
        stmt = (update(positions)
                .where(positions.c.order_no == bindparam("k_order"), positions.c.pos == bindparam("k_pos"))
                .values(manager=bindparam("manager")))
        conn.execute(stmt, params)


def sniff_source(data: bytes, filename: str) -> str | None:
    """Тип выгрузки по заголовкам — без разбора всей таблицы (CSV отчёта по отрезкам — десятки мегабайт)."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        for _title, headers in _sheet_headers(data):
            src = detect_source([str(h).strip() for h in headers])
            if src:
                return src
        return None
    head = _decode(data[:256 * 1024].split(b"\n", 1)[0])
    df = pd.read_csv(io.StringIO(head), dtype=str, sep=None, engine="python")
    return detect_source([str(c).strip() for c in df.columns])


def load_file(engine: Engine, source: str, data: bytes, filename: str) -> dict:
    if source not in LOADERS:
        raise ValueError("Неизвестный источник")
    df = read_table(data, filename)
    if source == "dispatcher" and filename.lower().endswith((".xlsx", ".xlsm")):
        return load_dispatcher(engine, df, filename, end_dates=_end_dates_sheet(data))
    return LOADERS[source](engine, df, filename)


def _end_dates_sheet(data: bytes) -> pd.DataFrame | None:
    for title, headers in _sheet_headers(data):
        normed = {p.norm_header(h) for h in headers}
        if {"заказ клиента", "дата конца"} <= normed and ("позиция заказа" in normed or "позиция заказа клиента" in normed) \
                and "рабочее место" not in normed:
            return pd.read_excel(io.BytesIO(data), sheet_name=title, dtype=object)
    return None


def detect_source(df_columns: list[str]) -> str | None:
    cols = set(df_columns)
    if set(SEGMENTS_REQUIRED) <= cols:
        return "segments"
    if set(SVETOFOR_REQUIRED) <= cols:
        return "svetofor"
    normed = {p.norm_header(c) for c in df_columns}
    if DISPATCHER_SIGNATURE <= normed:
        return "dispatcher"
    if PLAN_SIGNATURE <= normed:
        return "plan"
    return None
